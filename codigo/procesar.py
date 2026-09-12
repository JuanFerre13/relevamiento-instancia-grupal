# -*- coding: utf-8 -*-
"""
Etapa 2 — Procesamiento de la información.

Lee las respuestas del formulario (desde Google Sheets vía API, o desde un CSV
exportado) y genera:
  1. Un archivo XLSX con el formato "Relevamiento de Instancia grupal"
     (tabla de cantidades a la izquierda y de porcentajes a la derecha,
     con semáforo de colores).
  2. Un archivo JSON de resumen, que es el insumo de la etapa 3
     (generar_informe.py).

Uso típico:
  python procesar.py --sheet-id <ID_DE_LA_PLANILLA> --asistentes 15
  python procesar.py --csv respuestas.csv --unidad "Hospital Central" --asistentes 15

Si no se indica --unidad, se genera un archivo por cada unidad ejecutora
encontrada en las respuestas (aplicando los demás filtros).
"""

import argparse
import csv
import hashlib
import json
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from rutas import archivo_externo, recurso

# Rutas por defecto, resueltas contra la RAÍZ del proyecto (o contra el .exe) y
# no contra el directorio actual: config.json vive en `recursos/` y
# credenciales.json queda en la raíz, junto al ejecutable. La interfaz igual las
# pasa explícitas; esto es la red para el uso por consola desde cualquier lado.
CONFIG_POR_DEFECTO = str(recurso("recursos/config.json"))
CREDENCIALES_POR_DEFECTO = str(archivo_externo("credenciales.json"))

# ----------------------------------------------------------------------------
# Colores (semáforo) y estilos
# ----------------------------------------------------------------------------
COLOR_ROJO = "FF7C5C"      # En desacuerdo
COLOR_AMARILLO = "FFFF66"  # Neutro
COLOR_VERDE = "92D050"     # De acuerdo

FUENTE = "Arial"
F_NORMAL = Font(name=FUENTE, size=10)
F_NEGRITA = Font(name=FUENTE, size=10, bold=True)
F_TITULO = Font(name=FUENTE, size=11, bold=True)
F_CHICA = Font(name=FUENTE, size=8, bold=True)

BORDE_FINO = Border(*[Side(style="thin")] * 4)

FILL_ROJO = PatternFill("solid", fgColor=COLOR_ROJO)
FILL_AMARILLO = PatternFill("solid", fgColor=COLOR_AMARILLO)
FILL_VERDE = PatternFill("solid", fgColor=COLOR_VERDE)

CENTRADO = Alignment(horizontal="center", vertical="center", wrap_text=True)
IZQUIERDA = Alignment(horizontal="left", vertical="center", wrap_text=True)


# ----------------------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------------------
def normalizar(texto):
    """Minúsculas, sin tildes, espacios colapsados. Para comparar textos."""
    if texto is None:
        return ""
    texto = str(texto).strip().lower()
    texto = unicodedata.normalize("NFD", texto)
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", texto)


def buscar_columna(encabezados, nombre):
    """Devuelve el encabezado real que coincide (normalizado) con `nombre`."""
    objetivo = normalizar(nombre)
    for h in encabezados:
        if normalizar(h) == objetivo:
            return h
    # coincidencia parcial como último recurso (útil si Forms agrega espacios)
    for h in encabezados:
        if objetivo and objetivo in normalizar(h):
            return h
    return None


def es_afirmativa(valor, afirmativa="Sí"):
    """¿La celda responde que SÍ? Tolera cómo escriba la opción el formulario.

    Google Forms devuelve el texto literal de la opción elegida, que puede ser
    "Sí", "SI", "Si" o una frase que empieza así ("Sí, considero necesario…").
    Se compara normalizado (sin tildes ni mayúsculas) y se acepta el prefijo
    sólo cuando lo sigue un separador, para que "siempre" no cuente como "sí"."""
    v = normalizar(valor)
    objetivo = normalizar(afirmativa)
    if not v or not objetivo:
        return False
    if v == objetivo:
        return True
    return v.startswith(objetivo) and v[len(objetivo):len(objetivo) + 1] in " ,.;:-"


def detector_de_entrevistas(encabezados, config, log=print):
    """Arma el criterio de "esta fila pidió una entrevista individual".

    Es UNA sola definición porque de acá salen dos cosas que no se pueden
    contradecir: el número que va al informe firmado (`contar_entrevistas`) y la
    lista del botón «Ver entrevistas pendientes» (`entrevistas_solicitadas`).

    El criterio es acumulativo: hay que haber respondido que SÍ a la pregunta de
    confirmación Y haber dejado nombre y celular. Se eligió el más conservador
    de los posibles (2026-08-13, decisión del usuario): así ninguna entrevista
    contada en un informe firmado es incontactable, y el número no se infla con
    quien dejó sus datos sin querer una entrevista.

    Compatibilidad hacia atrás: si la columna de confirmación NO existe en la
    planilla, se vuelve al criterio viejo (nombre + celular). Eso mantiene
    andando los CSV y las planillas anteriores a que se agregara la pregunta.

    Devuelve `(predicado, col_nombre, col_celular, col_confirma)`; las columnas
    van para que quien llama pueda avisar qué falta."""
    cols = config["columnas"]
    nombre_col_nombre = cols.get("nombre_completo", "Nombre completo")
    nombre_col_celular = cols.get("celular", "Celular de contacto")
    nombre_col_confirma = cols.get("solicita_entrevista", "")
    afirmativa = cols.get("respuesta_afirmativa", "Sí")

    col_nombre = buscar_columna(encabezados, nombre_col_nombre)
    col_celular = buscar_columna(encabezados, nombre_col_celular)
    col_confirma = buscar_columna(encabezados, nombre_col_confirma) if nombre_col_confirma else None

    if nombre_col_confirma and not col_confirma:
        log("⚠ No aparece la columna de confirmación de entrevista "
            f'("{nombre_col_confirma}"). Se usa el criterio anterior: '
            "cuentan las filas con nombre y celular cargados.")

    def pidio_entrevista(r):
        if not (col_nombre and col_celular):
            return False
        tiene_datos = bool((r.get(col_nombre) or "").strip()
                           and (r.get(col_celular) or "").strip())
        if not tiene_datos:
            return False
        if not col_confirma:
            return True
        return es_afirmativa(r.get(col_confirma), afirmativa)

    return pidio_entrevista, col_nombre, col_celular, col_confirma


def nombre_archivo_seguro(texto):
    texto = normalizar(texto).replace(" ", "_")
    return re.sub(r"[^a-z0-9_]", "", texto) or "salida"


def _sin_tildes(texto):
    texto = unicodedata.normalize("NFD", str(texto or ""))
    return "".join(c for c in texto if unicodedata.category(c) != "Mn")


def fecha_mm_aa(fecha):
    """'mm-aa' a partir de una fecha tipo 20/7/26 o 20/07/2026 (mes y año, 2 dígitos)."""
    m = re.match(r"^\s*(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})\s*$", str(fecha or ""))
    if not m:
        return ""
    _dia, mes, anio = (int(x) for x in m.groups())
    return f"{mes:02d}-{anio % 100:02d}"


def nombre_carpeta_instancia(unidad, fecha):
    """Nombre de la subcarpeta de resultados: '{U.E.}-{mm}-{aa}'.

    La unidad va sin tildes ni espacios y conservando mayúsculas
    (p. ej. 'Centro Departamental Este' → 'CentroDepartamentalEste')."""
    unidad_limpia = re.sub(r"[^A-Za-z0-9]", "", _sin_tildes(unidad))
    mmaa = fecha_mm_aa(fecha)
    if not mmaa:  # fecha con formato no reconocido: fallback seguro sin romper
        mmaa = re.sub(r"[^A-Za-z0-9]", "", _sin_tildes(fecha))
    partes = [p for p in (unidad_limpia, mmaa) if p]
    return "-".join(partes) or "instancia"


# ----------------------------------------------------------------------------
# Lectura de datos
# ----------------------------------------------------------------------------
def _correo_cuenta_servicio(ruta_cred):
    try:
        return json.loads(Path(ruta_cred).read_text(encoding="utf-8")).get("client_email", "")
    except Exception:
        return ""


def _mensaje_error_google(error, worksheet, ruta_cred):
    """Traduce los errores de Google a algo que el cliente pueda entender.

    Sin esto, un ID mal copiado le muestra literalmente '<Response [404]>'."""
    texto = str(error)
    clase = type(error).__name__
    codigo = getattr(getattr(error, "response", None), "status_code", None)
    if codigo is None:
        for posible in (404, 403, 401):
            if str(posible) in texto:
                codigo = posible
                break

    correo = _correo_cuenta_servicio(ruta_cred)
    con_correo = f" El correo figura en credenciales.json: {correo}" if correo else ""

    if "WorksheetNotFound" in clase:
        return (f"La planilla no tiene ninguna hoja llamada «{worksheet}». "
                "Revisá el nombre en Configuración, o dejalo vacío para usar la primera hoja.")
    if codigo == 404 or "SpreadsheetNotFound" in clase or "not found" in texto.lower():
        return ("No se encontró ninguna planilla de Google con ese ID. Revisá el ID en "
                "Configuración: es la parte larga de la dirección web de la planilla, "
                "entre /d/ y /edit.")
    if codigo in (401, 403):
        return ("La planilla existe, pero el programa no tiene permiso para leerla. "
                "Abrila en Google, tocá «Compartir» y agregá como Lector el correo de la "
                "cuenta de servicio." + con_correo)
    if any(p in clase for p in ("Connection", "Timeout", "SSL")) or "connection" in texto.lower():
        return ("No se pudo conectar con Google. Revisá la conexión a internet y volvé a "
                "intentar. Si no hay conexión, podés usar el botón «Usar archivo descargado…».")
    return f"No se pudo leer la planilla de Google: {texto}"


def leer_google_sheets(sheet_id, worksheet, credenciales):
    try:
        import gspread
    except ImportError:
        raise RuntimeError("Falta la librería gspread. Ejecutá el instalador (instalar.bat) o: pip install -r requirements.txt")

    ruta_cred = Path(credenciales)
    if not ruta_cred.exists():
        raise RuntimeError(
            f"No se encontró el archivo de credenciales '{credenciales}'. "
            "Seguí los pasos de GUIA.md (sección 'Configurar la API de Google')."
        )
    try:
        cliente = gspread.service_account(filename=str(ruta_cred))
        planilla = cliente.open_by_key(sheet_id)
        hoja = planilla.worksheet(worksheet) if worksheet else planilla.sheet1
        filas = hoja.get_all_values()
    except Exception as e:
        raise RuntimeError(_mensaje_error_google(e, worksheet, ruta_cred)) from e

    if not filas:
        raise RuntimeError("La planilla está vacía.")
    encabezados = filas[0]
    registros = [dict(zip(encabezados, fila)) for fila in filas[1:] if any(c.strip() for c in fila)]
    return encabezados, registros


def leer_csv(ruta):
    with open(ruta, newline="", encoding="utf-8-sig") as f:
        lector = csv.DictReader(f)
        registros = [r for r in lector if any((v or "").strip() for v in r.values())]
        return lector.fieldnames, registros


# ----------------------------------------------------------------------------
# Procesamiento
# ----------------------------------------------------------------------------
def detectar_columnas_unidad(encabezados, cols_config):
    """Encuentra las columnas de Unidad Ejecutora aunque cambien de nombre.

    Toma toda columna cuyo encabezado empiece con "U.E" o "RAP" (p. ej.
    "U.E. Región Sur", "RAP NORTE"). Si no hay ninguna, usa la lista
    del config.json como estaba antes."""
    detectadas = []
    for h in encabezados:
        n = normalizar(h).replace(".", "").replace(" ", "")
        if n.startswith("ue") or n.startswith("rap"):
            detectadas.append(h)
    if detectadas:
        return detectadas
    return [c for c in (buscar_columna(encabezados, n) for n in cols_config) if c]


def unidad_de_registro(registro, columnas_reales):
    """Con el select dependiente, la U.E. queda en una sola de las columnas.

    Si la columna de U.E. dice solo "RAP" (valor intermedio del formulario),
    se sigue buscando el nombre concreto de la RAP en las demás columnas."""
    respaldo = ""
    for col in columnas_reales:
        valor = (registro.get(col) or "").strip()
        if not valor:
            continue
        if normalizar(valor) == "rap":
            respaldo = respaldo or valor
            continue
        return valor
    return respaldo


def solo_fecha(valor):
    """'28/7/2026 17:46:26' -> '28/7/2026'. La marca temporal de Google Forms
    trae la hora pegada y para comparar instancias sólo importa el día."""
    return str(valor or "").strip().split(" ")[0].strip()


def fecha_de_registro(registro, col_fecha, col_marca):
    """Fecha de la instancia. Usa la columna 'Fecha' del formulario y, si viene
    vacía, cae en la marca temporal que Google Forms completa sola.

    En la planilla real la pregunta de fecha quedó sin responder en todas las
    filas, así que sin este respaldo no se puede filtrar por fecha."""
    if col_fecha:
        valor = (registro.get(col_fecha) or "").strip()
        if valor:
            return solo_fecha(valor)
    if col_marca:
        return solo_fecha(registro.get(col_marca) or "")
    return ""


def region_de_registro(registro, col_region, columnas_unidad):
    """Región de la respuesta. Usa la columna 'Región' y, si está vacía, la
    deduce del nombre de la columna de U.E. que vino completa (los encabezados
    reales son del tipo 'U.E. Región Norte' / 'RAP ESTE')."""
    if col_region:
        valor = (registro.get(col_region) or "").strip()
        if valor:
            return valor
    for col in columnas_unidad:
        if not (registro.get(col) or "").strip():
            continue
        n = normalizar(col)
        for punto in ("norte", "sur", "este", "oeste"):
            if punto in n:
                return f"Región {punto.capitalize()}"
    return ""


def servicio_de_registro(registro, col_servicio):
    return (registro.get(col_servicio) or "").strip() if col_servicio else ""


def misma_fecha(a, b):
    """Compara fechas con tolerancia de formato: 18/7/2026 == 18/07/26."""
    na, nb = normalizar(a), normalizar(b)
    if na == nb:
        return True
    patron = r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})$"
    ma, mb = re.match(patron, na), re.match(patron, nb)
    if ma and mb:
        da, mma, ya = (int(x) for x in ma.groups())
        db, mmb, yb = (int(x) for x in mb.groups())
        return (da, mma, ya % 100) == (db, mmb, yb % 100)
    return False


def clasificar_respuesta(valor, opciones):
    v = normalizar(valor)
    if not v:
        return None  # omitida
    if v == normalizar(opciones["en_desacuerdo"]):
        return "en_desacuerdo"
    if v == normalizar(opciones["neutro"]):
        return "neutro"
    if v == normalizar(opciones["de_acuerdo"]):
        return "de_acuerdo"
    return None


def procesar_grupo(registros, config, encabezados):
    """Cuenta respuestas por pregunta y clasifica médicos / no médicos."""
    col_tipo = buscar_columna(encabezados, config["columnas"]["tipo_encuestado"])
    medicos = 0
    for r in registros:
        tipo = normalizar(r.get(col_tipo, "")) if col_tipo else ""
        if tipo == "medico":
            medicos += 1
    resumen = {
        "respuestas": len(registros),
        "medicos": medicos,
        "no_medicos": len(registros) - medicos,
        "secciones": [],
    }
    faltantes = []
    for seccion in config["secciones"]:
        datos_seccion = {"titulo": seccion["titulo"], "preguntas": []}
        for pregunta in seccion["preguntas"]:
            col = buscar_columna(encabezados, pregunta)
            if col is None:
                faltantes.append(pregunta)
                conteo = {"en_desacuerdo": 0, "neutro": 0, "de_acuerdo": 0}
            else:
                conteo = {"en_desacuerdo": 0, "neutro": 0, "de_acuerdo": 0}
                for r in registros:
                    clase = clasificar_respuesta(r.get(col, ""), config["opciones"])
                    if clase:
                        conteo[clase] += 1
            conteo["total"] = sum(conteo.values())
            datos_seccion["preguntas"].append({"texto": pregunta, **conteo})
        resumen["secciones"].append(datos_seccion)
    return resumen, faltantes


# ----------------------------------------------------------------------------
# Generación del XLSX
# ----------------------------------------------------------------------------
def escribir_encabezado_bloque(ws, fila, col_ini, info, config, ancho_bloque):
    """Título + datos generales de un bloque. Devuelve la fila siguiente."""
    c = lambda offset: get_column_letter(col_ini + offset)

    ws.merge_cells(f"{c(0)}1:{c(ancho_bloque - 1)}1")
    celda = ws.cell(row=1, column=col_ini, value=config["titulo"])
    celda.font = F_TITULO
    celda.alignment = CENTRADO

    filas_info = [
        ("U.E.", info["unidad"]),
        ("SERVICIO/ÁREA", info["servicio"]),
        ("FECHA", info["fecha"]),
        (None, None),
        ("Asistentes", info["asistentes"]),
        ("Respuestas", info["respuestas"]),
    ]
    f = 3
    for etiqueta, valor in filas_info:
        if etiqueta is None:
            f += 1
            continue
        ws.cell(row=f, column=col_ini, value=etiqueta).font = F_NEGRITA
        celda_v = ws.cell(row=f, column=col_ini + 1, value=valor)
        celda_v.font = F_NEGRITA if etiqueta in ("Respuestas",) else F_NORMAL
        if etiqueta == "Respuestas":
            info["celda_respuestas"] = f"${c(1)}${f}"
        f += 1

    # Médicos / No médicos con porcentaje calculado por fórmula
    ref = info["celda_respuestas"]
    for etiqueta, valor in (("Médicos", info["medicos"]), ("No Médicos", info["no_medicos"])):
        ws.cell(row=f, column=col_ini, value=etiqueta).font = F_NEGRITA
        ws.cell(row=f, column=col_ini + 1, value=valor).font = F_NORMAL
        celda_pct = ws.cell(row=f, column=col_ini + 2,
                            value=f"=IF({ref}=0,0,{c(1)}{f}/{ref})")
        celda_pct.number_format = "0 %"
        celda_pct.font = F_NORMAL
        f += 1

    # Entrevistas individuales del día (dato calculado; mismo valor en todas las
    # unidades de esa fecha). Va DESPUÉS de Médicos/No Médicos a propósito, para
    # no correr la celda "Respuestas" ($B$8) sobre la que se calculan los %.
    entrevistas = info.get("entrevistas_individuales")
    if entrevistas is not None:
        ws.cell(row=f, column=col_ini, value="Entrevistas individuales").font = F_NEGRITA
        ws.cell(row=f, column=col_ini + 1, value=entrevistas).font = F_NORMAL
        f += 1
    return f + 1


def escribir_bloque(ws, config, info, col_ini, modo, fila_inicio):
    """Escribe las secciones de preguntas. modo = 'cantidades' | 'porcentajes'."""
    ancho = 6  # n°, texto, en desacuerdo, neutro, de acuerdo, total
    ref_resp = info["celda_respuestas"]
    col_letra = lambda offset: get_column_letter(col_ini + offset)
    fila = fila_inicio

    for seccion in info["resumen"]["secciones"]:
        # Encabezado de sección
        ws.merge_cells(start_row=fila, start_column=col_ini,
                       end_row=fila, end_column=col_ini + 1)
        celda = ws.cell(row=fila, column=col_ini, value=seccion["titulo"])
        celda.font = F_NEGRITA
        celda.alignment = IZQUIERDA
        encabezados = [("En desacuerdo", FILL_ROJO), ("Neutro", FILL_AMARILLO),
                       ("De acuerdo", FILL_VERDE), ("Total", None)]
        for i, (texto, relleno) in enumerate(encabezados):
            cel = ws.cell(row=fila, column=col_ini + 2 + i, value=texto)
            cel.font = F_CHICA if texto != "Total" else F_NEGRITA
            cel.alignment = CENTRADO
            if relleno:
                cel.fill = relleno
            cel.border = BORDE_FINO
        ws.cell(row=fila, column=col_ini).border = BORDE_FINO
        ws.cell(row=fila, column=col_ini + 1).border = BORDE_FINO
        fila += 1

        for n, pregunta in enumerate(seccion["preguntas"], start=1):
            cel_n = ws.cell(row=fila, column=col_ini, value=n)
            cel_n.font = F_NORMAL
            cel_n.alignment = CENTRADO
            cel_t = ws.cell(row=fila, column=col_ini + 1, value=pregunta["texto"])
            cel_t.font = F_NORMAL
            cel_t.alignment = IZQUIERDA

            valores = [pregunta["en_desacuerdo"], pregunta["neutro"],
                       pregunta["de_acuerdo"], pregunta["total"]]
            rellenos = [FILL_ROJO, FILL_AMARILLO, FILL_VERDE, None]
            for i, (valor, relleno) in enumerate(zip(valores, rellenos)):
                col = col_ini + 2 + i
                if modo == "cantidades":
                    cel = ws.cell(row=fila, column=col,
                                  value=valor if valor else None)
                    if i == 3:  # Total = suma de las tres opciones
                        a, b = col_letra(2), col_letra(4)
                        cel.value = f"=SUM({a}{fila}:{b}{fila})"
                        cel.font = F_NEGRITA
                    else:
                        cel.font = F_NORMAL
                else:  # porcentajes: fórmula sobre la celda de cantidades
                    col_cant = get_column_letter(info["col_cant_ini"] + 2 + i)
                    cel = ws.cell(row=fila, column=col,
                                  value=f"=IF({ref_resp}=0,0,{col_cant}{fila}/{ref_resp})")
                    cel.number_format = "0 %"
                    cel.font = F_NEGRITA if i == 3 else F_NORMAL
                cel.alignment = CENTRADO
                if relleno:
                    cel.fill = relleno
                cel.border = BORDE_FINO
            cel_n.border = BORDE_FINO
            cel_t.border = BORDE_FINO
            fila += 1
        fila += 1  # fila en blanco entre secciones

    return fila


def generar_xlsx(info, config, ruta_salida):
    wb = Workbook()
    ws = wb.active
    ws.title = "Relevamiento"

    ANCHO_BLOQUE = 6
    COL_CANT = 1   # columna A
    COL_SEP = COL_CANT + ANCHO_BLOQUE          # columna G (separador)
    COL_PCT = COL_SEP + 1                      # columna H

    info["col_cant_ini"] = COL_CANT

    # Anchos de columna. openpyxl mide el ancho en "unidades de carácter";
    # las columnas A y H se piden en píxeles. Conversión aprox. (MDW=7):
    # px = round(width*7)+5  ->  width = (px-5)/7  (así el default 8.43 → 64 px).
    ANCHO_A_H_PX = 185
    ancho_a_h = (ANCHO_A_H_PX - 5) / 7  # ≈ 25.71 → ~185 px
    for col_ini in (COL_CANT, COL_PCT):
        ws.column_dimensions[get_column_letter(col_ini)].width = ancho_a_h
        ws.column_dimensions[get_column_letter(col_ini + 1)].width = 52
        for i in range(2, 5):
            ws.column_dimensions[get_column_letter(col_ini + i)].width = 12
        ws.column_dimensions[get_column_letter(col_ini + 5)].width = 8
    ws.column_dimensions[get_column_letter(COL_SEP)].width = 3

    # Encabezados de ambos bloques (la referencia a "Respuestas" queda en info)
    fila_inicio = escribir_encabezado_bloque(ws, 1, COL_CANT, info, config, ANCHO_BLOQUE)
    escribir_encabezado_bloque(ws, 1, COL_PCT, info, config, ANCHO_BLOQUE)

    # Los % de ambos bloques se calculan sobre la celda "Respuestas" del
    # bloque de cantidades (columna B, fila 8 según el encabezado de arriba).
    info["celda_respuestas"] = "$B$8"
    fila_fin = escribir_bloque(ws, config, info, COL_CANT, "cantidades", fila_inicio)
    escribir_bloque(ws, config, info, COL_PCT, "porcentajes", fila_inicio)

    # Nota al pie
    fila_nota = fila_fin
    nota = config.get("nota_pie")
    if nota:
        cel = ws.cell(row=fila_nota, column=COL_PCT + 1, value=nota)
        cel.font = F_NEGRITA
        fila_nota += 1

    # Alcance del informe (sólo en los agregados: qué fechas, unidades y
    # servicios se juntaron). Va al pie para no correr la celda $B$8.
    for texto in info.get("notas_alcance") or []:
        ws.cell(row=fila_nota, column=COL_PCT + 1, value=texto).font = F_NORMAL
        fila_nota += 1

    wb.save(ruta_salida)


# ----------------------------------------------------------------------------
# Tipos de informe
# ----------------------------------------------------------------------------
# Qué filtra cada uno:
#   instancia  → fecha + U.E. + servicio(s). Es el informe de siempre; si no se
#                elige U.E., genera uno por cada unidad con respuestas ese día
#                (la ventana ya no ofrece esa opción; queda para la consola).
#   multifecha → las fechas elegidas + U.E. + servicio(s), en UN solo informe.
#                Es UNA instancia repartida en dos o más jornadas (el mismo taller
#                dictado en días distintos, p. ej. para dos turnos del servicio).
#   unidad     → U.E. entera (todas las fechas y servicios), en UN solo informe.
# Entre los dos primeros no elige quien usa el programa: los separa
# `resolver_tipo()` según cuántas fechas se marcaron.
#
# Hubo un cuarto tipo, `region`, que juntaba todas las U.E. de una región en un
# solo informe. Se sacó el 2026-08-05 por pedido del usuario. La región de cada
# respuesta se sigue leyendo (`region_de_registro()` → `_region`): va como dato
# en el JSON de cualquier informe y la muestra `--listar`.
TIPOS_INFORME = {
    "instancia": "Instancia puntual",
    "multifecha": "Instancia fecha múltiple",
    "unidad": "Por Unidad Ejecutora",
}

# "multifecha" se llamaba "historico" y juntaba TODAS las fechas de la unidad, sin
# poder elegirlas. Se acepta el nombre viejo para no romper la línea de comandos.
ALIAS_TIPOS = {"historico": "multifecha"}


def normalizar_tipo(tipo):
    """Clave de TIPOS_INFORME, aceptando los nombres viejos de ALIAS_TIPOS."""
    tipo = ALIAS_TIPOS.get(tipo, tipo or "instancia")
    if tipo not in TIPOS_INFORME:
        raise RuntimeError(f"Tipo de informe desconocido: {tipo!r}.")
    return tipo


def resolver_tipo(tipo, fecha=None, fechas=None):
    """Decide entre `instancia` y `multifecha` según CUÁNTAS fechas se eligieron.

    La ventana tiene una sola opción "Instancia" y una lista de fechas para
    marcar: quien la usa no tiene por qué clasificar de antemano si su taller
    fue de un día o de varios. Acá se traduce esa elección al tipo real, y de
    eso dependen tres cosas que NO son cosméticas: el criterio de las
    entrevistas individuales (total del día vs. dentro del conjunto), el nombre
    de la carpeta y el bloque `<alcance_del_informe>` que se le manda a la IA.

    Con una sola fecha el resultado es idéntico al del informe puntual de
    siempre. Devuelve (tipo, fecha, fechas) ya coherentes entre sí."""
    tipo = normalizar_tipo(tipo)
    elegidas = [f for f in (fechas or []) if f]

    if tipo == "instancia" and len(elegidas) > 1:
        return "multifecha", None, elegidas
    if tipo == "multifecha" and len(elegidas) == 1:
        # Una sola jornada no es una instancia de fecha múltiple. Ojo: el
        # `historico` viejo SIN fechas elegidas sigue siendo multifecha (junta
        # todas), por eso se mira que haya exactamente una.
        return "instancia", elegidas[0], []
    if tipo == "instancia" and len(elegidas) == 1:
        return "instancia", fecha or elegidas[0], []
    return tipo, fecha, elegidas


def _clave_fecha(fecha):
    """Clave para ordenar fechas cronológicamente aunque vengan como texto."""
    m = re.match(r"^\s*(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})\s*$", str(fecha or ""))
    if not m:
        return (9999, 99, 99, str(fecha or ""))
    dia, mes, anio = (int(x) for x in m.groups())
    return (anio if anio > 99 else 2000 + anio, mes, dia, "")


def ordenar_fechas(fechas):
    return sorted({f for f in fechas if f}, key=_clave_fecha)


def texto_periodo(fechas, detallado=False):
    """Una sola fecha se muestra tal cual; varias, como rango.

    Con `detallado` (lo usa la instancia de fecha múltiple) las pocas fechas se
    enumeran en vez de mostrarse como período: son las jornadas de un mismo
    taller, y "del 28/7 al 4/8" haría pensar en algo continuo."""
    ordenadas = ordenar_fechas(fechas)
    if not ordenadas:
        return ""
    if len(ordenadas) == 1:
        return ordenadas[0]
    if detallado and len(ordenadas) <= 3:
        return " y ".join([", ".join(ordenadas[:-1]), ordenadas[-1]])
    return f"del {ordenadas[0]} al {ordenadas[-1]}"


def anotar_registros(registros, encabezados, config, log=print):
    """Agrega a cada fila los campos derivados `_unidad`, `_fecha`, `_region` y
    `_servicio`, y devuelve las columnas reales encontradas."""
    cols = config["columnas"]
    col_fecha = buscar_columna(encabezados, cols.get("fecha", "Fecha"))
    col_marca = buscar_columna(encabezados, cols.get("marca_temporal", "Marca temporal"))
    col_region = buscar_columna(encabezados, cols.get("region", "Región"))
    col_servicio = buscar_columna(encabezados, cols.get("servicio", "Servicio/Área"))
    cols_unidad = detectar_columnas_unidad(encabezados, cols.get("unidad", []))
    if not cols_unidad:
        log("⚠ No se detectaron columnas de Unidad Ejecutora (U.E/RAP) en la planilla.")
    for r in registros:
        r["_unidad"] = unidad_de_registro(r, cols_unidad)
        r["_fecha"] = fecha_de_registro(r, col_fecha, col_marca)
        r["_region"] = region_de_registro(r, col_region, cols_unidad)
        r["_servicio"] = servicio_de_registro(r, col_servicio)
    return {"fecha": col_fecha, "marca": col_marca, "region": col_region,
            "servicio": col_servicio, "unidad": cols_unidad}


def leer_respuestas(csv=None, sheet_id=None, worksheet=None,
                    credenciales=CREDENCIALES_POR_DEFECTO, config_path=CONFIG_POR_DEFECTO,
                    log=print):
    """Lee la planilla (o el CSV) y devuelve (config, encabezados, registros ya
    anotados). Lo comparten `procesar_datos` y `opciones_disponibles`."""
    with open(config_path, encoding="utf-8") as f:
        config = json.load(f)

    if csv:
        encabezados, registros = leer_csv(csv)
    else:
        if not sheet_id:
            raise RuntimeError("Falta el ID de la planilla de Google Sheets (configuralo primero).")
        encabezados, registros = leer_google_sheets(sheet_id, worksheet, credenciales)

    if not registros:
        raise RuntimeError("No se encontraron respuestas en la planilla.")

    anotar_registros(registros, encabezados, config, log=log)
    return config, encabezados, registros


def opciones_disponibles(csv=None, sheet_id=None, worksheet=None,
                         credenciales=CREDENCIALES_POR_DEFECTO, config_path=CONFIG_POR_DEFECTO,
                         log=print):
    """Qué hay realmente cargado en la planilla, para que la interfaz muestre
    listas para elegir en vez de campos de texto libre.

    Las fechas van de la más reciente a la más antigua (ver el comentario abajo).

    Los servicios vienen con su conteo y con el texto original, porque en el
    formulario es un campo libre y la misma área aparece escrita de varias
    formas ("Emergencia adultos", "Emergencia de adulto"…): la persona elige
    del listado las variantes que quiere juntar."""
    _config, _encabezados, registros = leer_respuestas(
        csv=csv, sheet_id=sheet_id, worksheet=worksheet,
        credenciales=credenciales, config_path=config_path, log=log)

    # Los servicios se agrupan por su forma normalizada, que es exactamente el
    # criterio con el que después se filtra: si en la planilla están
    # "Emergencia adultos" y "Emergencia Adultos", son el mismo y se muestran
    # una sola vez, para que el conteo del listado sea el que se va a obtener.
    grupos = {}
    for r in registros:
        valor = r["_servicio"]
        if not valor:
            continue
        clave = normalizar(valor)
        g = grupos.setdefault(clave, {"variantes": {}, "respuestas": 0})
        g["variantes"][valor] = g["variantes"].get(valor, 0) + 1
        g["respuestas"] += 1

    servicios = []
    for clave in sorted(grupos):
        g = grupos[clave]
        # Como etiqueta, la forma más usada (y a igual uso, la primera alfabética).
        etiqueta = sorted(g["variantes"], key=lambda v: (-g["variantes"][v], v))[0]
        servicios.append({
            "clave": clave,                     # forma normalizada, para cruzar
            "etiqueta": etiqueta,
            "respuestas": g["respuestas"],
            "variantes": sorted(g["variantes"]),
        })

    # Combinaciones que existen de verdad (fecha + unidad + región + servicio).
    # Con esto la interfaz encadena las listas y ofrece sólo lo que corresponde a
    # lo ya elegido: si el 28/7 en esa U.E. no hubo Laboratorio, no lo muestra.
    conteo_combinaciones = {}
    for r in registros:
        clave = (r["_fecha"], r["_unidad"], r["_region"], normalizar(r["_servicio"]))
        conteo_combinaciones[clave] = conteo_combinaciones.get(clave, 0) + 1
    combinaciones = [
        {"fecha": f, "unidad": u, "region": g, "servicio": s, "respuestas": n}
        for (f, u, g, s), n in sorted(conteo_combinaciones.items())
    ]

    unidades = {}
    for r in registros:
        if r["_unidad"]:
            unidades.setdefault(r["_unidad"], 0)
            unidades[r["_unidad"]] += 1

    return {
        "respuestas": len(registros),
        # De la más reciente a la más antigua: es el orden en que se busca, porque
        # el informe que se pide es casi siempre el de lo último relevado. Es sólo
        # el orden de las listas para elegir; donde importa para el cálculo
        # (`fechas_incluidas`, `periodo`, `texto_periodo`) se sigue usando el
        # cronológico de `ordenar_fechas()`.
        "fechas": list(reversed(ordenar_fechas(r["_fecha"] for r in registros))),
        "unidades": sorted(unidades),
        "unidades_conteo": unidades,
        "regiones": sorted({r["_region"] for r in registros if r["_region"]}),
        "servicios": servicios,
        "combinaciones": combinaciones,
    }


# ----------------------------------------------------------------------------
# Entrevistas individuales solicitadas
# ----------------------------------------------------------------------------
def clave_entrevista(unidad, fecha, nombre, contacto):
    """Identificador estable de un pedido de entrevista.

    La planilla no tiene un número de fila que se pueda usar: las respuestas se
    agregan todo el tiempo y la hoja se puede reordenar, así que la fila se
    identifica por su contenido. Se guarda el hash y no los datos, para que el
    archivo donde se anotan las entrevistas ya hechas no quede con nombres ni
    teléfonos adentro."""
    crudo = "|".join(normalizar(x) for x in (unidad, fecha, nombre, contacto))
    return hashlib.sha1(crudo.encode("utf-8")).hexdigest()[:16]


def entrevistas_solicitadas(csv=None, sheet_id=None, worksheet=None,
                            credenciales=CREDENCIALES_POR_DEFECTO, config_path=CONFIG_POR_DEFECTO,
                            log=print):
    """Quiénes pidieron una entrevista individual en el formulario.

    Son las filas que confirmaron que SÍ quieren una entrevista y además dejaron
    nombre y celular (`detector_de_entrevistas()`, el mismo criterio con el que
    se cuentan en los informes). Devuelve una lista de dicts con `clave`,
    `unidad`, `fecha`, `nombre`, `contacto` y `orden`, **de la más antigua a la
    más reciente**: es una cola de trabajo y lo que más esperó va primero.

    La fecha es la misma que usa todo el programa (`_fecha`): sale de la columna
    `Fecha` del formulario y, como en la planilla real esa columna viene vacía,
    en los hechos es el día de la marca temporal."""
    config, encabezados, registros = leer_respuestas(
        csv=csv, sheet_id=sheet_id, worksheet=worksheet,
        credenciales=credenciales, config_path=config_path, log=log)

    cols = config["columnas"]
    pidio_entrevista, col_nombre, col_celular, _col_confirma = detector_de_entrevistas(
        encabezados, config, log=log)
    if not (col_nombre and col_celular):
        faltan = [n for n, c in ((cols.get("nombre_completo", "Nombre completo"), col_nombre),
                                 (cols.get("celular", "Celular de contacto"), col_celular)) if not c]
        raise RuntimeError(
            "No se puede armar la lista de entrevistas: en la planilla no aparece(n) la(s) "
            "columna(s) " + " y ".join(f"«{n}»" for n in faltan) + ".")

    pedidos = []
    for r in registros:
        nombre = (r.get(col_nombre) or "").strip()
        contacto = (r.get(col_celular) or "").strip()
        if not pidio_entrevista(r):
            continue
        pedidos.append({
            "clave": clave_entrevista(r["_unidad"], r["_fecha"], nombre, contacto),
            "unidad": r["_unidad"],
            "fecha": r["_fecha"],
            "nombre": nombre,
            "contacto": contacto,
        })

    # Más ANTIGUAS arriba: es una cola de trabajo, y lo que más esperó es lo
    # primero que hay que atender. Dentro del mismo día, por nombre (el sort de
    # Python es estable, así que alcanza con ordenar dos veces).
    pedidos.sort(key=lambda p: normalizar(p["nombre"]))
    pedidos.sort(key=lambda p: _clave_fecha(p["fecha"]))

    # `orden` es el lugar que le toca a cada pedido en esta lista ya ordenada.
    # Viaja hasta el JS para que, al deshacer una marca, la fila vuelva a su
    # posición cronológica en vez de aterrizar arriba de todo. Se manda el índice
    # y no la fecha porque quien ordena es Python: el JS no parsea fechas.
    for i, p in enumerate(pedidos):
        p["orden"] = i
    return pedidos


def filtrar_registros(registros, tipo, fecha=None, unidad=None, servicios=None,
                      fechas=None):
    """Aplica los filtros que corresponden al tipo de informe elegido.

    `fecha` es la jornada única del informe puntual; `fechas` son las jornadas
    marcadas de una instancia de fecha múltiple (si no se marca ninguna, entran
    todas las de esa U.E. y servicio)."""
    seleccion = list(registros)
    servicios_norm = {normalizar(s) for s in (servicios or []) if s}
    fechas_elegidas = [f for f in (fechas or []) if f]

    if tipo == "instancia" and fecha:
        seleccion = [r for r in seleccion if misma_fecha(r["_fecha"], fecha)]
    if tipo == "multifecha" and fechas_elegidas:
        seleccion = [r for r in seleccion
                     if any(misma_fecha(r["_fecha"], f) for f in fechas_elegidas)]
    if tipo in ("instancia", "multifecha", "unidad") and unidad:
        seleccion = [r for r in seleccion if normalizar(r["_unidad"]) == normalizar(unidad)]
    if tipo in ("instancia", "multifecha") and servicios_norm:
        seleccion = [r for r in seleccion if normalizar(r["_servicio"]) in servicios_norm]
    return seleccion


def _error_sin_coincidencias(todos, tipo, fecha, unidad, servicios, fechas=None):
    """Mensaje en lenguaje simple que enumera lo que SÍ hay en la planilla."""
    def listar(valores, tope=15):
        vistos = sorted({v for v in valores if v})
        return ", ".join(vistos[:tope]) or "(ninguna)"

    partes = ["Ninguna respuesta coincide con lo que elegiste."]
    if (tipo == "instancia" and fecha) or (tipo == "multifecha" and fechas):
        partes.append("Fechas que hay en la planilla: "
                      + listar(r["_fecha"] for r in todos) + ".")
    if unidad:
        partes.append("Unidades que hay en la planilla: "
                      + listar(r["_unidad"] for r in todos) + ".")
    if servicios:
        partes.append("Servicios/Áreas que hay en la planilla: "
                      + listar(r["_servicio"] for r in todos) + ".")
    partes.append("Revisá que la combinación elegida tenga respuestas cargadas.")
    return RuntimeError(" ".join(partes))


def nombre_carpeta_salida(info):
    """Subcarpeta donde se guardan los archivos de un informe.

    La instancia puntual conserva el nombre de siempre ({U.E.}-{mm}-{aa}); los
    informes agregados llevan un sufijo que dice de qué tipo son, para que no
    se pisen entre ellos ni con los puntuales."""
    tipo = info.get("tipo", "instancia")
    limpio = lambda t: re.sub(r"[^A-Za-z0-9]", "", _sin_tildes(t or ""))

    if tipo == "instancia":
        return nombre_carpeta_instancia(info["unidad"], info["fecha"])

    unidades = info.get("unidades_incluidas") or []
    unidad = limpio(unidades[0] if unidades else info.get("unidad"))
    if tipo == "unidad":
        return f"{unidad}-Completo" if unidad else "Unidad-Completo"
    # Fecha múltiple: se agregan el servicio y el mes (o los meses) de las
    # jornadas, para que dos instancias de la misma U.E. y servicio no se pisen.
    # Se usa el nombre elegido para el informe si lo hay; si no, la primera variante.
    servicios = info.get("servicios_incluidos") or []
    nombre_servicio = info.get("servicio") or (servicios[0] if servicios else "")
    servicio = limpio(nombre_servicio)[:24]
    meses = [m for m in (fecha_mm_aa(f) for f in info.get("fechas_incluidas") or []) if m]
    periodo = ""
    if meses:
        periodo = meses[0] if meses[0] == meses[-1] else f"{meses[0]}a{meses[-1]}"
    partes = [p for p in (unidad, servicio, periodo, "VariasFechas") if p]
    return "-".join(partes)


def notas_de_alcance(info):
    """Renglones al pie del Excel que dejan por escrito qué abarca el informe.

    Van como nota y no como filas del encabezado a propósito: el bloque de
    arriba tiene posiciones fijas y los porcentajes se calculan sobre $B$8."""
    tipo = info.get("tipo", "instancia")
    if tipo == "instancia":
        return []

    notas = [f"Tipo de informe: {info.get('tipo_nombre', '')}."]
    fechas = info.get("fechas_incluidas") or []
    if fechas:
        # En la instancia de fecha múltiple las fechas son jornadas de un mismo
        # taller; en el informe de unidad, instancias distintas.
        etiqueta = "Jornadas de la instancia" if tipo == "multifecha" else "Instancias incluidas"
        notas.append(f"{etiqueta} ({len(fechas)}): {', '.join(fechas)}.")
    unidades = info.get("unidades_incluidas") or []
    if len(unidades) > 1:
        notas.append(f"Unidades Ejecutoras incluidas ({len(unidades)}): {', '.join(unidades)}.")
    servicios = info.get("servicios_incluidos") or []
    if servicios:
        notas.append(f"Servicios/Áreas incluidos ({len(servicios)}): {', '.join(servicios)}.")
    return notas


# ----------------------------------------------------------------------------
# Función principal reutilizable (la usa la app de escritorio y la consola)
# ----------------------------------------------------------------------------
def procesar_datos(csv=None, sheet_id=None, worksheet=None,
                   credenciales=CREDENCIALES_POR_DEFECTO, config_path=CONFIG_POR_DEFECTO,
                   tipo="instancia", unidad=None, fecha=None, servicios=None,
                   asistentes=0, servicio="", salida="salida",
                   fechas=None, log=print):
    """Procesa las respuestas y devuelve [(ruta_xlsx, ruta_json), ...].

    `tipo` es una de las claves de TIPOS_INFORME. Sólo el tipo "instancia"
    puede devolver más de un resultado (uno por unidad cuando no se elige una);
    los demás juntan todo en un único informe.

    `fecha` es la jornada del informe puntual y `fechas` las de una instancia de
    fecha múltiple; `resolver_tipo()` acomoda el tipo según cuántas se eligieron."""
    tipo, fecha, fechas = resolver_tipo(tipo, fecha, fechas)

    config, encabezados, registros = leer_respuestas(
        csv=csv, sheet_id=sheet_id, worksheet=worksheet,
        credenciales=credenciales, config_path=config_path, log=log)
    todos = list(registros)

    # Entrevistas individuales: filas que confirmaron que SÍ quieren una y además
    # dejaron nombre y celular (ver detector_de_entrevistas).
    cols_cfg = config["columnas"]
    pidio_entrevista, col_nombre, col_celular, _col_confirma = detector_de_entrevistas(
        encabezados, config, log=log)
    if not (col_nombre and col_celular):
        faltan_cols = []
        if not col_nombre:
            faltan_cols.append(cols_cfg.get("nombre_completo", "Nombre completo"))
        if not col_celular:
            faltan_cols.append(cols_cfg.get("celular", "Celular de contacto"))
        log("⚠ No se pudieron contar las entrevistas individuales: no aparece(n) la(s) columna(s) "
            + " y ".join(f'\"{c}\"' for c in faltan_cols) + " en la planilla. Se informará 0.")

    def contar_entrevistas(fechas_del_grupo, grupo):
        """En el informe de una instancia puntual es el total del DÍA (todas las
        unidades), que es como se viene informando. En los demás tipos —incluida
        la instancia de fecha múltiple— no sirve ese criterio, porque mezclaría
        instancias de otras unidades o regiones que cayeron el mismo día: se
        cuenta dentro del propio conjunto filtrado."""
        if not (col_nombre and col_celular):
            return 0
        if tipo == "instancia":
            fechas = {f for f in fechas_del_grupo if f}
            base = [r for r in todos
                    if any(misma_fecha(r["_fecha"], f) for f in fechas)] if fechas else todos
            return sum(1 for r in base if pidio_entrevista(r))
        return sum(1 for r in grupo if pidio_entrevista(r))

    seleccion = filtrar_registros(registros, tipo, fecha=fecha, unidad=unidad,
                                  servicios=servicios, fechas=fechas)
    if not seleccion:
        raise _error_sin_coincidencias(todos, tipo, fecha, unidad, servicios,
                                       fechas=fechas)

    # Sólo el informe de instancia puntual se abre en uno por unidad; el resto
    # junta todas las respuestas seleccionadas en un único informe.
    if tipo == "instancia":
        unidades = sorted({r["_unidad"] for r in seleccion if r["_unidad"]}) or [unidad or "Sin unidad"]
        grupos = [(u, [r for r in seleccion if r["_unidad"] == u]) for u in unidades]
    else:
        grupos = [(None, seleccion)]

    carpeta = Path(salida)
    carpeta.mkdir(parents=True, exist_ok=True)

    resultados = []
    avisos = set()
    for etiqueta_unidad, grupo in grupos:
        if not grupo:
            continue
        resumen, faltantes = procesar_grupo(grupo, config, encabezados)
        avisos.update(faltantes)

        fechas_grupo = ordenar_fechas(r["_fecha"] for r in grupo)
        unidades_grupo = sorted({r["_unidad"] for r in grupo if r["_unidad"]})
        servicios_grupo = sorted({r["_servicio"] for r in grupo if r["_servicio"]})
        if tipo == "instancia" and fecha:
            periodo = fecha
        else:
            periodo = texto_periodo(fechas_grupo, detallado=(tipo == "multifecha"))

        titulo_unidad = etiqueta_unidad or unidad or (unidades_grupo[0] if unidades_grupo else "Sin unidad")

        # El renglón SERVICIO/ÁREA del Excel y el título del informe.
        # `servicio` es el nombre que eligió la persona para el conjunto: manda
        # sobre la lista de variantes, porque juntar los textos del formulario da
        # cosas como "Laboratorio, Laboratorio/ pediatría", que en un informe
        # oficial parece un error de carga.
        if servicio:
            titulo_servicio = servicio
        elif servicios:
            titulo_servicio = ", ".join(servicios)
        elif tipo == "unidad":
            titulo_servicio = "Todos los servicios"
        else:
            titulo_servicio = ""

        info = {
            "tipo": tipo,
            "tipo_nombre": TIPOS_INFORME[tipo],
            "unidad": titulo_unidad,
            "servicio": titulo_servicio,
            "fecha": periodo,
            "periodo": periodo,
            "fechas_incluidas": fechas_grupo,
            "unidades_incluidas": unidades_grupo,
            "servicios_incluidos": servicios_grupo,
            # Dato informativo: de qué región es la unidad relevada.
            "region": grupo[0]["_region"] if grupo else "",
            "asistentes": asistentes or None,
            "respuestas": resumen["respuestas"],
            "medicos": resumen["medicos"],
            "no_medicos": resumen["no_medicos"],
            "entrevistas_individuales": contar_entrevistas(fechas_grupo, grupo),
            "resumen": resumen,
        }
        info["notas_alcance"] = notas_de_alcance(info)

        base = nombre_archivo_seguro(f"relevamiento_{titulo_unidad}_{periodo}")
        subcarpeta = carpeta / nombre_carpeta_salida(info)
        subcarpeta.mkdir(parents=True, exist_ok=True)
        ruta_xlsx = subcarpeta / f"{base}.xlsx"
        ruta_json = subcarpeta / f"{base}.json"

        generar_xlsx(info, config, ruta_xlsx)

        json_info = {k: v for k, v in info.items()
                     if k not in ("resumen", "col_cant_ini", "celda_respuestas")}
        json_info["secciones"] = resumen["secciones"]
        total = resumen["respuestas"] or 1
        for seccion in json_info["secciones"]:
            for p in seccion["preguntas"]:
                p["pct_en_desacuerdo"] = round(100 * p["en_desacuerdo"] / total)
                p["pct_neutro"] = round(100 * p["neutro"] / total)
                p["pct_de_acuerdo"] = round(100 * p["de_acuerdo"] / total)
                p["pct_total"] = round(100 * p["total"] / total)
        with open(ruta_json, "w", encoding="utf-8") as f:
            json.dump(json_info, f, ensure_ascii=False, indent=2)

        log(f"✔ {titulo_unidad}: {resumen['respuestas']} respuestas → {subcarpeta.name}/{ruta_xlsx.name}")
        resultados.append((str(ruta_xlsx), str(ruta_json)))

    if avisos:
        log("⚠ Preguntas del config.json que no aparecen como columnas en la planilla:")
        for p in sorted(avisos):
            log(f"   - {p}")

    return resultados


def _a_entero(valor):
    """Devuelve int si el valor es numérico (o texto de dígitos); si no, None."""
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return int(valor)
    if isinstance(valor, str) and valor.strip().lstrip("-").isdigit():
        return int(valor.strip())
    return None


def _valor_celda_ods(celda, ns):
    """Valor de una celda de LibreOffice Calc: número si lo es, si no el texto."""
    tipo = celda.get(f"{{{ns['office']}}}value-type")
    if tipo in ("float", "percentage", "currency"):
        crudo = celda.get(f"{{{ns['office']}}}value")
        try:
            numero = float(crudo)
        except (TypeError, ValueError):
            return None
        return int(numero) if numero == int(numero) else numero
    texto = "\n".join("".join(p.itertext()) for p in celda.findall("text:p", ns)).strip()
    return texto or None


def _filas_ods(ruta, columnas=6):
    """Lee una planilla .ods (LibreOffice Calc) sin librerías extra: es un ZIP
    con un content.xml adentro.

    Se saltean las filas vacías: el resto del código busca las etiquetas y las
    preguntas por su contenido, no por el número de fila. Eso además evita las
    decenas de miles de filas vacías que Calc declara al final con el atributo
    `number-rows-repeated`."""
    import xml.etree.ElementTree as ET
    import zipfile

    ns = {
        "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
        "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
        "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    }
    try:
        with zipfile.ZipFile(ruta) as z:
            arbol = ET.fromstring(z.read("content.xml"))
    except Exception as e:
        raise RuntimeError(
            "No se pudo abrir la planilla de LibreOffice. Puede estar dañada o "
            f"no ser un archivo .ods válido. ({e})"
        )

    tabla = arbol.find(f".//{{{ns['table']}}}table")
    if tabla is None:
        raise RuntimeError("La planilla de LibreOffice no tiene ninguna hoja de cálculo.")

    filas = []
    for tr in tabla.findall(f"{{{ns['table']}}}table-row"):
        valores = []
        for tc in tr.findall(f"{{{ns['table']}}}table-cell"):
            repetidas = int(tc.get(f"{{{ns['table']}}}number-columns-repeated", 1) or 1)
            valor = _valor_celda_ods(tc, ns)
            valores.extend([valor] * min(repetidas, columnas))
            if len(valores) >= columnas:
                break
        if any(v is not None for v in valores):
            filas.append((valores + [None] * columnas)[:columnas])
    return filas


def _filas_xlsx(ruta, columnas=6):
    wb = load_workbook(ruta, data_only=True)
    ws = wb.active
    return [[c.value for c in fila] for fila in ws.iter_rows(min_row=1, max_col=columnas)]


FORMATOS_PLANILLA = {".xlsx": _filas_xlsx, ".xlsm": _filas_xlsx, ".ods": _filas_ods}


def filas_de_planilla(ruta, columnas=6):
    """Filas de la planilla, venga de Excel (.xlsx) o de LibreOffice (.ods)."""
    extension = Path(ruta).suffix.lower()
    lector = FORMATOS_PLANILLA.get(extension)
    if lector:
        return lector(ruta, columnas)
    if extension in (".odt", ".doc", ".docx"):
        raise RuntimeError(
            f"El archivo elegido es un documento de texto ({extension}), no una planilla. "
            "Este botón necesita la planilla del relevamiento, la de las tablas con los "
            "números: en LibreOffice Calc se guarda como .ods y en Excel como .xlsx."
        )
    raise RuntimeError(
        f"No se puede leer un archivo {extension or 'sin extensión'}. La planilla tiene que "
        "estar guardada como .ods (LibreOffice Calc) o .xlsx (Excel)."
    )


def datos_desde_xlsx(ruta_xlsx, config_path=CONFIG_POR_DEFECTO):
    """Lee un XLSX en el formato oficial del relevamiento y reconstruye el mismo
    resumen (dict) que produce procesar_datos, para poder generar el informe a
    partir de una planilla hecha a mano.

    Se apoya en la estructura del bloque de cantidades (columnas A–F): las filas
    de encabezado tienen una etiqueta de texto en A (U.E., FECHA, Respuestas…) y
    su valor en B; las filas de pregunta tienen el número en A, el texto en B y
    los conteos En desacuerdo/Neutro/De acuerdo en C/D/E.

    Devuelve (datos, preguntas_no_encontradas)."""
    with open(config_path, encoding="utf-8") as f:
        config = json.load(f)

    etiquetas = {}   # etiqueta normalizada (col A) -> valor (col B)
    conteos = {}     # texto de pregunta normalizado -> (en_desacuerdo, neutro, de_acuerdo)
    for fila in filas_de_planilla(ruta_xlsx):
        a, b = fila[0], fila[1]
        a_num = _a_entero(a)
        if a_num is None and isinstance(a, str) and a.strip():
            etiquetas.setdefault(normalizar(a), b)  # fila de encabezado o título de sección
        elif a_num is not None and isinstance(b, str) and b.strip():
            # Fila de pregunta (número en A, texto en B). En el formato oficial un
            # conteo de 0 se guarda como celda vacía, así que None se toma como 0.
            ed = _a_entero(fila[2]) or 0
            ne = _a_entero(fila[3]) or 0
            da = _a_entero(fila[4]) or 0
            conteos.setdefault(normalizar(b), (ed, ne, da))

    def etiqueta(*nombres):
        for n in nombres:
            v = etiquetas.get(normalizar(n))
            if v is not None:
                return v
        return None

    respuestas = _a_entero(etiqueta("Respuestas")) or 0

    secciones = []
    faltantes = []
    for seccion in config["secciones"]:
        datos_sec = {"titulo": seccion["titulo"], "preguntas": []}
        for pregunta in seccion["preguntas"]:
            key = normalizar(pregunta)
            conteo = conteos.get(key)
            if conteo is None:  # coincidencia parcial de respaldo
                for k, v in conteos.items():
                    if key and (key in k or k in key):
                        conteo = v
                        break
            if conteo is None:
                faltantes.append(pregunta)
                ed, ne, da = 0, 0, 0
            else:
                ed, ne, da = conteo
            total = ed + ne + da
            den = respuestas or 1
            datos_sec["preguntas"].append({
                "texto": pregunta,
                "en_desacuerdo": ed, "neutro": ne, "de_acuerdo": da, "total": total,
                "pct_en_desacuerdo": round(100 * ed / den),
                "pct_neutro": round(100 * ne / den),
                "pct_de_acuerdo": round(100 * da / den),
                "pct_total": round(100 * total / den),
            })
        secciones.append(datos_sec)

    total_preguntas = sum(len(s["preguntas"]) for s in config["secciones"])
    if faltantes and len(faltantes) == total_preguntas:
        raise RuntimeError(
            "No se reconoció ninguna pregunta en la planilla. Asegurate de que sea una "
            "planilla en el formato oficial del relevamiento (con la tabla de cantidades "
            "en las columnas A a F, y los textos de las preguntas igual que en el formulario)."
        )

    datos = {
        "tipo": "instancia",
        "tipo_nombre": TIPOS_INFORME["instancia"],
        "unidad": str(etiqueta("U.E.", "U.E", "Unidad") or "").strip() or "Sin unidad",
        "servicio": str(etiqueta("SERVICIO/ÁREA", "Servicio/Área", "Servicio") or ""),
        "fecha": str(etiqueta("FECHA", "Fecha") or "").strip(),
        "asistentes": _a_entero(etiqueta("Asistentes")),
        "respuestas": respuestas,
        "medicos": _a_entero(etiqueta("Médicos")) or 0,
        "no_medicos": _a_entero(etiqueta("No Médicos")) or 0,
        "entrevistas_individuales": _a_entero(etiqueta("Entrevistas individuales")) or 0,
        "secciones": secciones,
    }
    return datos, faltantes


def _consola_utf8():
    """La consola de Windows suele usar cp1252 y rompe al imprimir ✔/⚠.
    Forzamos UTF-8 en la salida cuando se puede (no aplica si no hay consola)."""
    for flujo in (sys.stdout, sys.stderr):
        try:
            flujo.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass


def main():
    _consola_utf8()
    parser = argparse.ArgumentParser(description="Procesa las respuestas del formulario y genera el XLSX + JSON de resumen.")
    origen = parser.add_mutually_exclusive_group(required=True)
    origen.add_argument("--sheet-id", help="ID de la planilla de Google Sheets (está en la URL)")
    origen.add_argument("--csv", help="Ruta a un CSV exportado (alternativa sin API)")
    parser.add_argument("--worksheet", help="Nombre de la hoja dentro de la planilla (por defecto, la primera)")
    parser.add_argument("--credenciales", default=CREDENCIALES_POR_DEFECTO, help="Archivo JSON de la cuenta de servicio de Google")
    parser.add_argument("--config", default=CONFIG_POR_DEFECTO, help="Archivo de configuración (secciones y preguntas)")
    parser.add_argument("--tipo", default="instancia",
                        choices=sorted(list(TIPOS_INFORME) + list(ALIAS_TIPOS)),
                        help="Tipo de informe: instancia (fecha+U.E.+servicio), multifecha "
                             "(varias jornadas de una misma instancia) o unidad (U.E. entera)")
    parser.add_argument("--unidad", help="Filtrar por una Unidad Ejecutora específica")
    parser.add_argument("--fecha", help="Filtrar por fecha (solo con --tipo instancia)")
    parser.add_argument("--fechas", nargs="+", metavar="FECHA",
                        help="Jornadas de la instancia (solo con --tipo multifecha). Sin "
                             "esto entran todas las fechas de esa U.E. y servicio.")
    parser.add_argument("--servicios", nargs="+", metavar="SERVICIO",
                        help="Uno o más valores de Servicio/Área a incluir (se pueden juntar "
                             "las variantes de un mismo servicio)")
    parser.add_argument("--asistentes", type=int, default=0, help="Cantidad de asistentes a la instancia (dato manual)")
    parser.add_argument("--servicio", default="", help="Texto para el campo SERVICIO/ÁREA")
    parser.add_argument("--salida", default="salida", help="Carpeta de salida")
    parser.add_argument("--listar", action="store_true",
                        help="No genera nada: muestra las fechas, unidades, servicios y "
                             "regiones que hay cargados en la planilla")
    parser.add_argument("--entrevistas", action="store_true",
                        help="No genera nada: muestra quiénes pidieron una entrevista "
                             "individual (todas, sin filtrar por las ya realizadas: eso lo "
                             "lleva la aplicación)")
    args = parser.parse_args()

    try:
        if args.entrevistas:
            pedidos = entrevistas_solicitadas(
                csv=args.csv, sheet_id=args.sheet_id, worksheet=args.worksheet,
                credenciales=args.credenciales, config_path=args.config)
            print(f"Entrevistas individuales solicitadas: {len(pedidos)}\n")
            for p in pedidos:
                print(f"   {p['fecha']:>12}  {p['unidad']}")
                print(f"                 {p['nombre']} — {p['contacto']}")
            return

        if args.listar:
            opciones = opciones_disponibles(
                csv=args.csv, sheet_id=args.sheet_id, worksheet=args.worksheet,
                credenciales=args.credenciales, config_path=args.config)
            print(f"Respuestas: {opciones['respuestas']}")
            print(f"\nFechas ({len(opciones['fechas'])}):")
            for f in opciones["fechas"]:
                print(f"   - {f}")
            print(f"\nUnidades Ejecutoras ({len(opciones['unidades'])}):")
            for u in opciones["unidades"]:
                print(f"   - {u}  ({opciones['unidades_conteo'][u]} respuestas)")
            print(f"\nRegiones ({len(opciones['regiones'])}):")
            for r in opciones["regiones"]:
                print(f"   - {r}")
            print(f"\nServicios/Áreas ({len(opciones['servicios'])}):")
            for s in opciones["servicios"]:
                print(f"   - {s['etiqueta']}  ({s['respuestas']} respuestas)")
                if len(s["variantes"]) > 1:
                    print(f"       variantes: {', '.join(s['variantes'])}")
            return

        procesar_datos(csv=args.csv, sheet_id=args.sheet_id, worksheet=args.worksheet,
                       credenciales=args.credenciales, config_path=args.config,
                       tipo=args.tipo, unidad=args.unidad, fecha=args.fecha,
                       fechas=args.fechas, servicios=args.servicios,
                       asistentes=args.asistentes, servicio=args.servicio,
                       salida=args.salida)
    except RuntimeError as e:
        sys.exit(f"Error: {e}")


if __name__ == "__main__":
    main()
