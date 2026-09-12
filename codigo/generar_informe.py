# -*- coding: utf-8 -*-
"""
Etapa 3 — Generación del informe con IA (API de Claude).

Toma el JSON de resumen generado por procesar.py, le agrega como contexto los
informes de ejemplo que estén en la carpeta `informes_ejemplo/` (.txt, .md o
.docx) y le pide a Claude que redacte un informe nuevo con el mismo estilo.

Uso:
  set ANTHROPIC_API_KEY=sk-ant-...        (Windows)
  export ANTHROPIC_API_KEY=sk-ant-...     (Linux/Mac)
  python generar_informe.py salida/relevamiento_centro_auxiliar_101026.json

El informe se guarda en la misma carpeta como .md y .docx.
"""

import argparse
import json
import os
import re
import sys
from pathlib import Path

from rutas import recurso

MODELO = "claude-opus-5"

# OJO al cambiar de modelo desde acá o desde ajustes.json: NO se manda el
# parámetro `thinking` a propósito. Cada modelo tiene su propio valor por
# defecto (Opus 5 razona solo; Opus 4.8 y anteriores no) y las formas de
# configurarlo son incompatibles entre sí, así que no mandarlo es lo único que
# funciona con cualquiera de los dos sin tocar código.

# Tope de tokens de la respuesta. El informe más largo del corpus real usa unos
# 2.500, y los que juntan varias instancias (fecha múltiple, unidad) son más largos, así que
# se deja margen amplio: sólo se paga por lo que el modelo escribe de verdad, no
# por el tope. Antes estaba en 4.000 y el informe podía cortarse a la mitad.
# Con Opus 5 el tope cubre además lo que el modelo razona antes de escribir.
MAX_TOKENS = 16000

# Precio del modelo, por millón de tokens (dólares). Está acá, en un solo lugar,
# nada más que para estimar el costo en el registro. OJO: las tarifas cambian;
# verificar en https://docs.claude.com/en/docs/about-claude/models
PRECIO_ENTRADA_POR_MILLON = 5.0
PRECIO_SALIDA_POR_MILLON = 25.0

# Recursos de solo lectura: se resuelven a la carpeta del proyecto como script,
# o a la carpeta temporal embebida cuando corre como .exe (PyInstaller).
CARPETA_EJEMPLOS = str(recurso("recursos/informes_ejemplo"))
PLANTILLA = str(recurso("recursos/plantilla.docx"))
ARCHIVO_ESTILO = str(recurso("recursos/estilo_informe.md"))

INSTRUCCIONES = """Sos un/a analista de la Unidad de Mediación y Convivencia que redacta \
informes de relevamiento de instancias grupales para unidades ejecutoras del sector salud.

Te voy a dar:
1. Uno o más informes anteriores como EJEMPLO de estilo, estructura, tono y extensión.
2. Los datos procesados (cantidades y porcentajes de respuesta por pregunta, agrupadas \
en secciones: Comunicación, Relaciones Interpersonales, Trabajo en Equipo, Motivación y \
Valoración de la Instancia).

El campo "tipo" de los datos dice qué informe hay que redactar: "instancia" es el taller \
puntual de un día (el caso de los ejemplos); "multifecha" es UNA sola instancia cuyos \
talleres se dictaron en dos o más jornadas; "unidad" resume VARIAS instancias de una misma \
Unidad Ejecutora. Si aparece un bloque <alcance_del_informe>, leelo antes que nada y \
respetá lo que dice.

Redactá un informe NUEVO sobre los datos proporcionados, imitando fielmente el estilo \
de los ejemplos. Reglas:
- Usá únicamente los datos proporcionados; no inventes cifras ni información.
- Destacá fortalezas (indicadores con alto % de acuerdo) y puntos de atención \
(indicadores con mayor % en desacuerdo o neutro).
- Si en los ejemplos hay secciones fijas (introducción, metodología, conclusiones, \
recomendaciones), reproducí esa misma estructura.
- Escribí en español, en el registro formal-institucional de los ejemplos.
- Si se incluye una <guia_de_estilo>, seguila al pie de la letra: define la \
estructura, las secciones y el tono exactos del informe.
- Devolvé SOLO el informe en formato Markdown, sin comentarios adicionales. \
Usá # para el encabezado del informe, ## para los títulos de sección, \
**negrita** y *cursiva* como indica la guía. No incluyas encabezado ni pie \
de página institucional: eso lo aplica la plantilla automáticamente.
- NO uses notas al pie ni la sintaxis [^1]: el informe se convierte a Word y esas \
marcas saldrían escritas tal cual en el documento firmado. Si hace falta remitir al \
cuadro de datos, hacelo dentro de la propia oración (por ejemplo "según surge del \
cuadro de relevamiento adjunto")."""


def instrucciones_completas():
    base = INSTRUCCIONES
    ruta = Path(ARCHIVO_ESTILO)
    if ruta.exists():
        base += "\n\n<guia_de_estilo>\n" + ruta.read_text(encoding="utf-8") + "\n</guia_de_estilo>"
    return base


def _filas_a_markdown(filas):
    """Filas de una tabla (listas de texto) → tabla Markdown de pipes.

    Se emite en el mismo formato que `markdown_a_docx()` sabe leer de vuelta
    (fila de celdas + separadora `|---|---|`), así el ejemplo que ve el modelo
    está escrito tal como se espera que él escriba las suyas. Lo comparten el
    lector de .docx y el de .odt."""
    limpias = []
    for fila in filas:
        # Una celda puede tener varios párrafos: se aplasta a una línea, porque
        # un salto adentro de la fila partiría la tabla Markdown en dos.
        celdas = [" ".join(celda.split()) for celda in fila]
        if any(celdas):
            limpias.append(celdas)
    if not limpias:
        return ""
    # El ancho lo fija la primera fila (el encabezado). Las demás se rellenan o
    # se recortan a esa medida: una fila con otra cantidad de celdas —pasa en
    # .odt, donde las vacías del final no se escriben— daría una tabla torcida.
    columnas = len(limpias[0])
    lineas = ["| " + " | ".join((f + [""] * columnas)[:columnas]) + " |" for f in limpias]
    separadora = "|" + "|".join(["---"] * columnas) + "|"
    return "\n".join([lineas[0], separadora] + lineas[1:])


def _tabla_a_markdown(tabla):
    """Tabla de python-docx → tabla Markdown de pipes."""
    return _filas_a_markdown([[c.text for c in fila.cells] for fila in tabla.rows])


def _texto_de_docx(archivo):
    """Texto de un .docx en el orden del documento, tablas incluidas.

    `doc.paragraphs` NO devuelve los párrafos que están dentro de celdas de
    tabla, así que las dinámicas de columnas («¿Qué voy a hacer más? / qué
    diferente / qué dejar de hacer») desaparecían de los ejemplos: el modelo
    veía informes sin una sola tabla mientras `estilo_informe.md` le pedía
    producirlas. Se recorre el cuerpo elemento por elemento en vez de usar
    `doc.paragraphs` + `doc.tables` para no perder el orden: una tabla tiene que
    quedar donde está en el informe, no amontonada al final."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(archivo))
    partes = []
    for elemento in doc.element.body.iterchildren():
        if elemento.tag == qn("w:p"):
            texto = Paragraph(elemento, doc).text.strip()
        elif elemento.tag == qn("w:tbl"):
            texto = _tabla_a_markdown(Table(elemento, doc))
        else:
            continue          # secciones, marcadores, etc.: no aportan texto
        if texto:
            partes.append(texto)
    return "\n".join(partes)


def _texto_de_odt(archivo):
    """Texto de un documento de LibreOffice Writer (.odt), tablas incluidas.

    El cliente escribe los informes en LibreOffice, así que los ejemplos llegan
    en .odt y no en .docx. No agrega ninguna dependencia: un .odt es un ZIP con
    un `content.xml` adentro, igual que el .ods que ya lee `procesar.py`, y se
    parsea con `zipfile` + `xml.etree` de la biblioteca estándar.

    Se recorre el cuerpo en orden y se despacha por tag, con el mismo criterio
    que `_texto_de_docx()`: los párrafos como texto plano y las tablas como
    Markdown de pipes, en el lugar donde están. El membrete no molesta porque en
    OpenDocument el encabezado y el pie viven en `styles.xml`, no acá."""
    import xml.etree.ElementTree as ET
    import zipfile

    ns_office = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
    ns_text = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
    ns_table = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"

    try:
        with zipfile.ZipFile(archivo) as z:
            arbol = ET.fromstring(z.read("content.xml"))
    except Exception as e:
        raise RuntimeError(
            "No se pudo abrir el documento de LibreOffice. Puede estar dañado o "
            f"no ser un archivo .odt válido. ({e})")

    cuerpo = arbol.find(f".//{{{ns_office}}}text")
    if cuerpo is None:
        raise RuntimeError("El documento de LibreOffice no tiene texto.")

    # Las notas al pie se quitan del árbol antes de leer nada. En OpenDocument
    # van INLINE dentro del párrafo, así que `itertext()` las metía en el medio
    # de la frase: "…surgen como más relevantes los siguientes datos1Ver cuadro
    # adjunto anexo.:" — así estaban 11 de los 12 ejemplos .odt. Tres razones
    # para descartarlas y no reubicarlas: (1) no son contenido del informe, son
    # punteros al Excel adjunto ("Ver Cuadro Relevamiento de Intervención
    # Grupal"); (2) el lector de .docx tampoco las ve, porque ahí viven en
    # footnotes.xml y no en el cuerpo, así que sacarlas deja los dos formatos
    # con el mismo criterio; (3) `markdown_a_docx()` no sabe generar notas al
    # pie, así que mostrárselas al modelo era enseñarle a escribir algo que no
    # se puede convertir — escribía "[^1]" y salía crudo en el informe firmado.
    for padre in arbol.iter():
        for hijo in list(padre):
            if hijo.tag != f"{{{ns_text}}}note":
                continue
            if hijo.tail:   # el texto pegado después de la nota sí es del informe
                hermanos = list(padre)
                anterior = hermanos[hermanos.index(hijo) - 1] if hermanos.index(hijo) else None
                if anterior is not None:
                    anterior.tail = (anterior.tail or "") + hijo.tail
                else:
                    padre.text = (padre.text or "") + hijo.tail
            padre.remove(hijo)

    def texto_plano(elemento):
        """Todo el texto de adentro, incluidos los `text:span` con que Writer
        parte un párrafo en cada cambio de formato."""
        return " ".join("".join(elemento.itertext()).split())

    def celdas_de(fila):
        celdas = []
        for tc in fila:
            if tc.tag not in (f"{{{ns_table}}}table-cell",
                              f"{{{ns_table}}}covered-table-cell"):
                continue
            repetidas = int(tc.get(f"{{{ns_table}}}number-columns-repeated", 1) or 1)
            celdas.extend([texto_plano(tc)] * min(repetidas, 20))
        return celdas

    partes = []

    def recorrer(elemento):
        for hijo in elemento:
            if hijo.tag in (f"{{{ns_text}}}p", f"{{{ns_text}}}h"):
                texto = texto_plano(hijo)
                if texto:
                    partes.append(texto)
            elif hijo.tag in (f"{{{ns_text}}}list", f"{{{ns_text}}}list-item",
                              f"{{{ns_text}}}section"):
                recorrer(hijo)   # viñetas y secciones: importa el texto de adentro
            elif hijo.tag == f"{{{ns_table}}}table":
                # Con iter() y no findall() para tomar también las filas que
                # Writer envuelve en `table:table-header-rows`, sin perder el orden.
                filas = [celdas_de(tr) for tr in hijo.iter(f"{{{ns_table}}}table-row")]
                tabla = _filas_a_markdown(filas)
                if tabla:
                    partes.append(tabla)

    recorrer(cuerpo)
    return "\n".join(partes)


def cargar_ejemplos(carpeta):
    ejemplos = []
    ruta = Path(carpeta)
    if not ruta.exists():
        return ejemplos
    for archivo in sorted(ruta.iterdir()):
        ext = archivo.suffix.lower()
        if archivo.stem.lower() in ("leeme", "readme", "léeme"):
            continue
        try:
            if ext in (".txt", ".md"):
                ejemplos.append((archivo.name, archivo.read_text(encoding="utf-8")))
            elif ext == ".docx":
                ejemplos.append((archivo.name, _texto_de_docx(archivo)))
            elif ext == ".odt":
                ejemplos.append((archivo.name, _texto_de_odt(archivo)))
        except Exception as e:
            print(f"⚠ No se pudo leer el ejemplo {archivo.name}: {e}")
    return ejemplos


def _sin_notas_al_pie(markdown):
    """Saca la sintaxis de notas al pie, que el Word generado no sabe hacer.

    `markdown_a_docx()` no genera notas al pie de verdad (python-docx no tiene
    API para eso y la plantilla no define ningún estilo, ver la regla de arriba),
    así que un `[^1]` terminaba escrito tal cual en el informe firmado: la misma
    fuga que tenían los `|` de las tablas antes de que se parsearan. A los
    ejemplos ya se les quitan las notas y a la IA se le pide que no las use (ver
    INSTRUCCIONES); esto es la red para cuando las use igual.

    El marcador se saca de la oración y el texto de la nota se conserva como un
    párrafo «Nota: …» al final, que es donde una nota al pie se termina leyendo.
    Se conserva y no se descarta porque suele ser la remisión al cuadro de datos
    ("Ver Cuadro Relevamiento de Intervención Grupal")."""
    notas = []

    def guardar(m):
        notas.append(m.group(2).strip())
        return ""

    # Definición en su propia línea: "[^1]: Ver Cuadro Relevamiento…"
    texto = re.sub(r"^\[\^([^\]]+)\]:[ \t]*(.*)$", guardar, markdown, flags=re.M)
    # Marcador dentro de la oración: "…los siguientes datos[^1]:"
    texto = re.sub(r"\[\^[^\]]+\]", "", texto)
    if notas:
        texto = texto.rstrip() + "\n\n" + "\n".join(f"Nota: {n}" for n in notas if n)
    return texto


def markdown_a_docx(markdown, ruta_docx, titulo=None, plantilla=None):
    """Markdown → Word con la estética institucional.

    Si existe `plantilla.docx` (hoja membretada), el informe se genera sobre
    ella, conservando su encabezado y pie de página en todas las hojas.
    Estilo del cuerpo: Calibri 11, texto justificado, títulos en negrita
    subrayada (nivel 1 centrado), como los informes de la Unidad."""
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ruta_plantilla = Path(plantilla) if plantilla else Path(PLANTILLA)
    if ruta_plantilla.exists():
        doc = Document(str(ruta_plantilla))
        for p in list(doc.paragraphs):  # limpiar el cuerpo; el membrete queda
            el = p._element
            el.getparent().remove(el)
    else:
        doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    def agregar_texto(parrafo, texto, subrayado=False, negrita=None, tam=None):
        partes = re.split(r"(\*\*.+?\*\*|\*[^*]+\*)", texto)
        for parte in partes:
            if not parte:
                continue
            run = parrafo.add_run()
            if parte.startswith("**") and parte.endswith("**"):
                run.text = parte[2:-2]
                run.bold = True
            elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
                run.text = parte[1:-1]
                run.italic = True
            else:
                run.text = parte
            if negrita is not None:
                run.bold = negrita
            if subrayado:
                run.underline = True
            if tam:
                run.font.size = Pt(tam)

    def parrafo_lista(estilo, marca):
        """Ítem de lista. Devuelve (párrafo, prefijo).

        `plantilla.docx` no define los estilos de lista de Word ("List Bullet" /
        "List Number"), así que cuando faltan se arma la viñeta a mano con
        sangría, que además es como se ven en los informes del corpus. Sin este
        respaldo, cualquier informe con viñetas no llegaba a generar el .docx."""
        try:
            return doc.add_paragraph(style=estilo), ""
        except KeyError:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.space_after = Pt(2)
            return p, marca

    def celdas_de(linea):
        """Celdas de una fila markdown: | a | b | -> ['a', 'b']."""
        return [c.strip() for c in linea.strip().strip("|").split("|")]

    def agregar_tabla(filas):
        """Tabla de Word con bordes dibujados a mano.

        `plantilla.docx` NO trae ningún estilo de tabla (ni siquiera "Table
        Grid"), así que pedirle un estilo lanzaría KeyError y, como esto corre
        dentro de un try en generar(), el .docx no se generaría y el cliente se
        quedaría sólo con el .md. Por eso los bordes se ponen por XML: no
        dependen de que la plantilla tenga nada definido."""
        tabla = doc.add_table(rows=len(filas), cols=len(filas[0]))
        bordes = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borde = OxmlElement(f"w:{lado}")
            borde.set(qn("w:val"), "single")
            borde.set(qn("w:sz"), "4")
            borde.set(qn("w:color"), "808080")
            bordes.append(borde)
        tabla._element.tblPr.append(bordes)

        for i, fila in enumerate(filas):
            for j, celda in enumerate(fila[:len(filas[0])]):
                parrafo = tabla.cell(i, j).paragraphs[0]
                # La primera fila es el encabezado de la dinámica: va en negrita.
                agregar_texto(parrafo, celda, negrita=True if i == 0 else None)
        doc.add_paragraph()   # aire después de la tabla
        return tabla

    lineas = _sin_notas_al_pie(markdown).splitlines()
    indice = 0
    while indice < len(lineas):
        linea = lineas[indice].rstrip()
        indice += 1
        if not linea.strip():
            continue

        # ¿Arranca una tabla? Fila de celdas seguida de la fila separadora
        # (|---|---|). La guía de estilo las pide para las dinámicas de columnas.
        siguiente = lineas[indice].strip() if indice < len(lineas) else ""
        if (linea.strip().startswith("|")
                and re.match(r"^\|[\s:\-|]+\|$", siguiente)):
            filas = [celdas_de(linea)]
            indice += 1                       # saltear la fila separadora
            while indice < len(lineas) and lineas[indice].strip().startswith("|"):
                filas.append(celdas_de(lineas[indice]))
                indice += 1
            agregar_tabla(filas)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)", linea)
        if m:
            nivel = len(m.group(1))
            p = doc.add_paragraph()
            if nivel == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            agregar_texto(p, m.group(2), subrayado=True, negrita=True,
                          tam=12 if nivel <= 2 else 11)
        elif re.match(r"^\s*[-*]\s+", linea):
            p, prefijo = parrafo_lista("List Bullet", "• ")
            agregar_texto(p, prefijo + re.sub(r"^\s*[-*]\s+", "", linea))
        elif re.match(r"^\s*\d+[.)]\s+", linea):
            numero = re.match(r"^\s*(\d+)[.)]\s+", linea).group(1)
            p, prefijo = parrafo_lista("List Number", f"{numero}. ")
            agregar_texto(p, prefijo + re.sub(r"^\s*\d+[.)]\s+", "", linea))
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            agregar_texto(p, linea)
    doc.save(str(ruta_docx))


def _bloque_alcance(datos):
    """Aviso explícito de qué abarca el informe cuando NO es de una instancia
    puntual. Los ejemplos de estilo son todos de un taller de un día, así que
    sin esto el modelo redacta 'la instancia del día de hoy' aunque los datos
    resuman meses o varias unidades."""
    tipo = datos.get("tipo", "instancia")
    if tipo == "instancia":
        return None

    fechas = datos.get("fechas_incluidas") or []
    unidades = datos.get("unidades_incluidas") or []
    # OJO: acá va el nombre ELEGIDO para el informe, no `servicios_incluidos`.
    # Esa lista son las formas en que se tipeó el área en el formulario
    # ("Puerta", "Puerta emergencia"…) y es registro interno: al mandarla, el
    # modelo escribía "en sus distintas denominaciones de Puerta y Emergencia de
    # adultos" en un informe firmado (reportado el 2026-08-05).
    servicio = (datos.get("servicio") or "").strip()

    if tipo == "multifecha":
        # No es un agregado de instancias distintas: es una sola instancia que se
        # dictó en varios días. Si no se le dice así, el modelo la redacta como un
        # seguimiento en el tiempo.
        lineas = [
            "ATENCIÓN: este informe corresponde a UNA sola instancia cuyos talleres se "
            "dictaron en más de una jornada. Es un único informe, no un seguimiento en "
            "el tiempo: seguí las reglas del tipo «multifecha» de la <guia_de_estilo>.",
            f"Alcance: {len(fechas)} jornada(s) de la misma instancia ({datos.get('periodo', '')}).",
        ]
        if fechas:
            lineas.append(f"Jornadas en que se dictó: {', '.join(fechas)}.")
    else:
        lineas = [
            f"ATENCIÓN: este NO es el informe de un taller puntual. Es un informe de tipo "
            f"«{datos.get('tipo_nombre', tipo)}» y seguí las reglas de la <guia_de_estilo> "
            f"para ese tipo (encabezado, Metodología y foco del Análisis).",
            f"Alcance: {len(fechas)} instancia(s) relevada(s) en el período {datos.get('periodo', '')}.",
        ]
        if fechas:
            lineas.append(f"Fechas comprendidas: {', '.join(fechas)}.")
    if unidades:
        lineas.append(f"Unidades Ejecutoras comprendidas ({len(unidades)}): {', '.join(unidades)}.")
    if servicio and tipo != "unidad":
        lineas.append(
            f"Servicio/Área: «{servicio}». Ese es el nombre del servicio y la ÚNICA forma de "
            "nombrarlo en todo el informe. No menciones ninguna otra denominación, ni aclares "
            "que el área figura escrita de varias maneras.")
        lineas.append(
            f"El informe se limita a «{servicio}» dentro de esa Unidad Ejecutora: no digas "
            "que abarca la totalidad de lo relevado en la unidad.")
    lineas.append(
        "Los conteos y porcentajes vienen SUMADOS sobre todas esas respuestas: no están "
        "abiertos por fecha ni por unidad. Por lo tanto no afirmes que algo mejoró o empeoró "
        "en el tiempo, ni atribuyas un porcentaje a una unidad o a una jornada en particular."
    )
    return "<alcance_del_informe>\n" + "\n".join(lineas) + "\n</alcance_del_informe>"


def _datos_para_la_ia(datos):
    """El JSON tal como se le muestra al modelo.

    Le saca `servicios_incluidos` cuando el informe es de un servicio concreto
    (instancia o fecha múltiple): esa lista son las formas en que se tipeó el
    área en el formulario libre —"Puerta", "Puerta emergencia", "Emergencia de
    adultos"— y existe como registro de qué respuestas entraron, no como
    información del informe. Teniéndola a la vista, el modelo la usaba: salió
    firmado un "abarcando las áreas de Emergencia (en sus distintas
    denominaciones de Puerta y Emergencia de adultos)" cuando el nombre elegido
    para el servicio era "Puerta de Emergencia" (2026-08-05). El nombre correcto
    ya viaja en el campo `servicio`.

    En el informe de unidad sí se conserva: ahí enumerar los servicios
    comprendidos es parte de lo que se pide. El archivo .json en disco queda
    intacto en los dos casos.

    También le saca `notas_alcance`, que es el texto al pie del Excel y trae la
    misma lista cruda por otra puerta (era la segunda vía de la fuga). Todo lo
    que dice ya está en el bloque <alcance_del_informe> y en los otros campos."""
    fuera = {"notas_alcance"}
    if datos.get("tipo") != "unidad":
        fuera.add("servicios_incluidos")
    return {k: v for k, v in datos.items() if k not in fuera}


def _armar_contenido(datos, lista_ejemplos, indicaciones=None, entrevistas_confidenciales=None):
    """Arma el pedido y lo devuelve partido en (estable, variable).

    La parte ESTABLE son los informes de ejemplo: son idénticos en cada llamada y
    pesan ~7.500 tokens, así que van primero y se cachean. La parte VARIABLE es
    todo lo que cambia de un informe a otro. El corte está acá y no en otro lado
    porque el caché de la API es por prefijo: alcanza con que cambie un byte
    antes del punto de corte para perder todo lo que viene después."""
    estables = []
    for nombre, texto in lista_ejemplos:
        estables.append(f"<informe_ejemplo nombre=\"{nombre}\">\n{texto}\n</informe_ejemplo>")

    bloques = []
    alcance = _bloque_alcance(datos)
    if alcance:
        bloques.append(alcance)
    bloques.append(f"<datos_nueva_instancia>\n"
                   f"{json.dumps(_datos_para_la_ia(datos), ensure_ascii=False, indent=2)}\n"
                   f"</datos_nueva_instancia>")
    if datos.get("medicos") or datos.get("no_medicos"):
        # Pedido expreso del usuario (2026-08-06): el desglose médico / no médico
        # tiene que figurar en la Metodología. Va acá y no sólo en
        # `estilo_informe.md` porque el corpus empuja en contra: NINGUNO de los 16
        # informes de ejemplo lo trae, así que una regla suelta en la guía pierde
        # contra 16 ejemplos que lo omiten (es lo que ya pasó con la sección de
        # entrevistas, prohibida dos veces en la guía y escrita igual). Por eso
        # el bloque nombra los números y dice explícitamente que los ejemplos no
        # lo traen y que hay que incluirlo igual.
        medicos = datos.get("medicos") or 0
        no_medicos = datos.get("no_medicos") or 0
        formularios = medicos + no_medicos
        aviso = (
            f"Composición del personal que respondió el formulario: de los {formularios} "
            f"formularios recibidos, {medicos} corresponden a personal médico y {no_medicos} a "
            "personal no médico. Incluí ese desglose en la Metodología de trabajo, a continuación "
            "de la cantidad de participantes («…de los cuales X corresponden al personal médico y "
            "Z al personal no médico»). Los informes de ejemplo NO traen este desglose: incluilo "
            "igual, es un pedido expreso de la Unidad."
        )
        # `asistentes` es un dato manual y puede ser mayor que las respuestas (no
        # todos los que concurren completan el formulario). Sin esta aclaración el
        # desglose se pega al número de concurrentes y la cuenta no cierra.
        if datos.get("asistentes") and datos["asistentes"] != formularios:
            aviso += (
                f" ATENCIÓN: concurrieron {datos['asistentes']} personas pero se recibieron "
                f"{formularios} formularios. El desglose es de los FORMULARIOS, no de los "
                f"asistentes: no lo presentes como la composición de las {datos['asistentes']} "
                "personas que concurrieron, porque no cerraría la cuenta."
            )
        bloques.append(aviso)
    if indicaciones:
        bloques.append(f"Indicaciones adicionales: {indicaciones}")
    if entrevistas_confidenciales:
        bloques.append(
            "El siguiente material proviene de entrevistas individuales confidenciales realizadas "
            "en el marco de la instancia. Incorporalo como una sección propia del informe, ubicada "
            "DESPUÉS del Análisis (luego de «D- Motivación») y ANTES de «Conclusiones y recomendaciones». "
            "La sección debe COMENZAR exactamente con la frase: «Es dable señalar que se solicitaron X "
            "entrevistas Individuales Confidenciales, de las cuales surge:», reemplazando X por el valor "
            "de \"entrevistas_individuales\" del JSON. A continuación desarrollá en prosa formal los "
            "conceptos aportados, cuidando la confidencialidad (sin nombres ni datos identificatorios):\n"
            "<entrevistas_individuales_confidenciales>\n"
            + entrevistas_confidenciales
            + "\n</entrevistas_individuales_confidenciales>"
        )
    elif datos.get("entrevistas_individuales"):
        # Sin material aportado el modelo escribía IGUAL la sección entera,
        # inventando qué habían dicho las personas en las entrevistas: "surge la
        # existencia de dificultades vinculares que se vienen sosteniendo en el
        # tiempo…". Medido el 2026-08-06 con los dos corpus de ejemplos, o sea
        # que no depende de cuántos ejemplos haya. `estilo_informe.md` ya lo
        # prohíbe dos veces (la regla de la sección y "no inventar testimonios")
        # y no alcanzó: la presión de ver el número en el JSON y la sección en
        # los informes de ejemplo le gana a la guía. Por eso la prohibición va
        # acá, pegada a los datos, y nombra lo que NO hay que hacer.
        # Va en la parte VARIABLE porque depende del JSON de cada instancia.
        bloques.append(
            "NO se aportó material de entrevistas individuales confidenciales. El dato "
            f"\"entrevistas_individuales\": {datos['entrevistas_individuales']} es solamente la "
            "CANTIDAD de entrevistas que se solicitaron: mencionala en la Metodología, como indica "
            "la guía, y nada más.\n"
            "Está PROHIBIDO incluir la sección que empieza «Es dable señalar que se solicitaron X "
            "entrevistas Individuales Confidenciales, de las cuales surge:», y está PROHIBIDO "
            "afirmar, resumir, insinuar o desarrollar qué se dijo en esas entrevistas: no tenés esa "
            "información. No la deduzcas de los porcentajes del formulario ni la tomes de los "
            "informes de ejemplo, que sí traen esa sección porque a ellos se les aportó el material."
        )
    bloques.append("Redactá ahora el informe de la nueva instancia.")
    return "\n\n".join(estables), "\n\n".join(bloques)


def exportar_prompt(json_resumen, ejemplos=CARPETA_EJEMPLOS, indicaciones=None,
                    entrevistas_confidenciales=None, log=print):
    """Modo sin clave: crea un .txt listo para pegar en claude.ai (gratis)."""
    ruta_json = Path(json_resumen)
    with open(ruta_json, encoding="utf-8") as f:
        datos = json.load(f)
    lista_ejemplos = cargar_ejemplos(ejemplos)
    if not lista_ejemplos:
        log(f"⚠ No hay informes en '{ejemplos}/'. El texto se generará sin ejemplos de estilo.")
    estable, variable = _armar_contenido(
        datos, lista_ejemplos, indicaciones, entrevistas_confidenciales)
    # Para pegar en claude.ai va todo junto: el corte sólo le sirve al caché de la API.
    contenido = "\n\n".join(p for p in (instrucciones_completas(), estable, variable) if p)
    ruta = Path(f"{ruta_json.with_suffix('')}_para_claude.txt")
    ruta.write_text(contenido, encoding="utf-8")
    log(f"✔ Texto para claude.ai creado: {ruta.name}")
    return str(ruta)


def _registrar_consumo(respuesta, log):
    """Deja en el registro cuántos tokens costó el informe y cuánto salió.

    El conteo de caché sirve además para confirmar que el caché está pegando:
    a partir del segundo informe de una tanda, `cache_read` tiene que ser > 0."""
    uso = getattr(respuesta, "usage", None)
    if uso is None:
        return
    entrada = getattr(uso, "input_tokens", 0) or 0
    salida = getattr(uso, "output_tokens", 0) or 0
    cache_leido = getattr(uso, "cache_read_input_tokens", 0) or 0
    cache_escrito = getattr(uso, "cache_creation_input_tokens", 0) or 0

    # Leer del caché sale ~0,1x y escribirlo ~1,25x respecto del precio de entrada.
    costo = (
        (entrada + cache_leido * 0.1 + cache_escrito * 1.25) * PRECIO_ENTRADA_POR_MILLON
        + salida * PRECIO_SALIDA_POR_MILLON
    ) / 1_000_000

    detalle = f"entrada {entrada}, salida {salida}"
    if cache_leido or cache_escrito:
        detalle += f", caché leído {cache_leido}, caché escrito {cache_escrito}"
    log(f"   Consumo: {detalle} tokens · aprox. US$ {costo:.3f}")


def generar(json_resumen, ejemplos=CARPETA_EJEMPLOS, modelo=MODELO,
            indicaciones=None, entrevistas_confidenciales=None, log=print):
    """Genera el informe con la API de Claude. Devuelve (ruta_md, ruta_docx|None)."""
    if not os.environ.get("ANTHROPIC_API_KEY"):
        raise RuntimeError(
            "Falta la clave de la API de Claude (ANTHROPIC_API_KEY). "
            "Cargala en la configuración de la aplicación."
        )
    try:
        import anthropic
    except ImportError:
        raise RuntimeError("Falta la librería anthropic. Ejecutá el instalador (instalar.bat) o: pip install -r requirements.txt")

    ruta_json = Path(json_resumen)
    with open(ruta_json, encoding="utf-8") as f:
        datos = json.load(f)

    lista_ejemplos = cargar_ejemplos(ejemplos)
    if not lista_ejemplos:
        log(f"⚠ No hay informes en '{ejemplos}/'. Se generará sin ejemplos de estilo.")

    estable, variable = _armar_contenido(
        datos, lista_ejemplos, indicaciones, entrevistas_confidenciales)

    # El mensaje va en dos bloques con el punto de corte del caché entre medio.
    # Como el `system` se manda antes que los `messages`, ese único corte cachea
    # las instrucciones + la guía de estilo + los informes de ejemplo, que son
    # el 86% del pedido y no cambian nunca.
    partes = []
    if estable:
        partes.append({"type": "text", "text": estable,
                       "cache_control": {"type": "ephemeral"}})
    partes.append({"type": "text", "text": variable})

    cliente = anthropic.Anthropic()
    log(f"Generando informe con IA ({modelo})… puede demorar un minuto.")
    respuesta = cliente.messages.create(
        model=modelo,
        max_tokens=MAX_TOKENS,
        system=instrucciones_completas(),
        messages=[{"role": "user", "content": partes}],
    )

    _registrar_consumo(respuesta, log)

    # El modelo puede declinar la solicitud: en ese caso no hay informe que guardar.
    if respuesta.stop_reason == "refusal":
        raise RuntimeError(
            "La IA no quiso redactar este informe (declinó la solicitud). "
            "Suele pasar si el texto de antecedentes o de entrevistas incluye algo "
            "que el modelo interpreta como sensible. Revisá esos campos y reintentá."
        )

    informe = "".join(b.text for b in respuesta.content if b.type == "text").strip()
    if informe.startswith("```"):
        informe = re.sub(r"^```[a-z]*\n|\n```$", "", informe)

    if not informe:
        raise RuntimeError(
            "La IA devolvió una respuesta vacía. Volvé a intentar; si se repite, "
            "puede haber un problema con la clave de la API o con el modelo elegido."
        )

    # Si el modelo llegó al tope, el texto quedó cortado a mitad de frase. Se
    # avisa fuerte pero se guardan igual los archivos: un informe cortado sirve
    # más que ninguno, siempre que se sepa que lo está.
    if respuesta.stop_reason == "max_tokens":
        log("⚠ ATENCIÓN: el informe quedó CORTADO — la IA llegó al límite de largo.")
        log("   Revisá el final del documento antes de usarlo y volvé a generarlo.")

    base = ruta_json.with_suffix("")
    ruta_md = Path(f"{base}_informe.md")
    ruta_docx = Path(f"{base}_informe.docx")
    ruta_md.write_text(informe, encoding="utf-8")
    try:
        markdown_a_docx(informe, ruta_docx)
        log(f"✔ Informe generado: {ruta_docx.name}")
        return str(ruta_md), str(ruta_docx)
    except Exception as e:
        log(f"✔ Informe generado: {ruta_md.name} (no se pudo crear el .docx: {e})")
        return str(ruta_md), None


def _consola_utf8():
    """La consola de Windows suele usar cp1252 y rompe al imprimir ✔/⚠.
    Forzamos UTF-8 en la salida cuando se puede (no aplica si no hay consola)."""
    for flujo in (sys.stdout, sys.stderr):
        try:
            flujo.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass


def main():
    _consola_utf8()
    parser = argparse.ArgumentParser(description="Genera el informe con la API de Claude a partir del JSON de resumen.")
    parser.add_argument("json_resumen", nargs="?",
                        help="Ruta al .json generado por procesar.py (o usá --desde-xlsx)")
    parser.add_argument("--desde-xlsx", metavar="PLANILLA",
                        help="Genera el informe a partir de una planilla en formato oficial hecha "
                             "a mano: .ods (LibreOffice Calc) o .xlsx (Excel). Reconstruye el JSON "
                             "de resumen y lo usa como insumo.")
    parser.add_argument("--ejemplos", default=CARPETA_EJEMPLOS, help="Carpeta con informes de ejemplo")
    parser.add_argument("--modelo", default=MODELO, help="Modelo de Claude a utilizar")
    parser.add_argument("--indicaciones", help="Indicaciones extra para el informe (opcional)")
    parser.add_argument("--entrevistas-confidenciales",
                        help="Texto con lo que surge de las entrevistas individuales confidenciales "
                             "(se agrega como sección tras el Análisis y antes de las Conclusiones)")
    parser.add_argument("--sin-clave", action="store_true",
                        help="No usa la API: crea un .txt para pegar en claude.ai")
    args = parser.parse_args()

    if args.desde_xlsx:
        import procesar
        datos, faltantes = procesar.datos_desde_xlsx(args.desde_xlsx)
        if faltantes:
            print(f"⚠ {len(faltantes)} pregunta(s) no se encontraron en la planilla (van en 0).")
        args.json_resumen = str(Path(args.desde_xlsx).with_suffix("")) + "_resumen.json"
        with open(args.json_resumen, "w", encoding="utf-8") as f:
            json.dump(datos, f, ensure_ascii=False, indent=2)
    if not args.json_resumen:
        parser.error("indicá un archivo .json o usá --desde-xlsx <archivo.xlsx>")

    try:
        if args.sin_clave:
            exportar_prompt(args.json_resumen, args.ejemplos, args.indicaciones,
                            args.entrevistas_confidenciales)
        else:
            generar(args.json_resumen, args.ejemplos, args.modelo, args.indicaciones,
                    args.entrevistas_confidenciales)
    except RuntimeError as e:
        sys.exit(f"Error: {e}")


if __name__ == "__main__":
    main()
