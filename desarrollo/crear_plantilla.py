# -*- coding: utf-8 -*-
r"""
Genera `recursos\plantilla.docx`: la hoja membretada sobre la que se escriben
todos los informes Word.

`markdown_a_docx()` abre esta plantilla, **borra los párrafos del cuerpo y
conserva encabezado y pie**, así que lo único que tiene que traer es el membrete
y los ajustes de página. El cuerpo lo escribe el programa.

En una instalación real este archivo se reemplaza por la papelería de la
organización (normalmente un .docx que provee el área de comunicación, con los
logos en el encabezado). El que genera este script es **ficticio**: existe para
que el repositorio funcione de punta a punta sin publicar papelería ajena.

Uso:  python desarrollo\crear_plantilla.py [--sobrescribir]
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
DESTINO = RAIZ / "recursos" / "plantilla.docx"

ORGANIZACION = "SERVICIO DE SALUD"
UNIDAD = "Unidad de Mediación y Convivencia"
PIE = "Unidad de Mediación y Convivencia · Av. Ejemplo 1234 · Tel. (000) 000 0000 · convivencia@ejemplo.org"

# Azul institucional del proyecto (el mismo del ícono y de la interfaz).
AZUL = RGBColor(0x1F, 0x4E, 0x79)

# Estilos que hay que SACAR de la plantilla. No es un capricho: la papelería
# real tampoco los define, y `markdown_a_docx()` tiene dos caminos alternativos
# escritos justamente para eso (`parrafo_lista()` arma la viñeta a mano con
# sangría, `agregar_tabla()` dibuja los bordes por XML). Una plantilla nacida de
# un `Document()` en blanco SÍ los trae, y esos caminos dejarían de ejercitarse:
# el repositorio pasaría a probar un escenario que en producción no ocurre.
ESTILOS_A_QUITAR = ("List Bullet", "List Number", "Table Grid")


def _sin_bordes(tabla):
    """Saca los bordes de la tabla del membrete: es maquetación, no un cuadro."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borde = OxmlElement(f"w:{lado}")
        borde.set(qn("w:val"), "none")
        bordes.append(borde)
    tabla._element.tblPr.append(bordes)


def construir():
    doc = Document()

    # Página A4 con los márgenes de la papelería institucional.
    seccion = doc.sections[0]
    seccion.page_width, seccion.page_height = Cm(21), Cm(29.7)
    seccion.top_margin = seccion.bottom_margin = Cm(2.5)
    seccion.left_margin = seccion.right_margin = Cm(3)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # ENCABEZADO. Va en una tabla de dos columnas, como la papelería real: a la
    # izquierda queda el lugar del logo (acá vacío, sin imagen) y a la derecha
    # el nombre de la organización y de la unidad.
    encabezado = seccion.header
    tabla = encabezado.add_table(rows=1, cols=2, width=Cm(15))
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    _sin_bordes(tabla)

    hueco = tabla.cell(0, 0).paragraphs[0]
    corrida = hueco.add_run("[ logo ]")
    corrida.font.size = Pt(9)
    corrida.font.color.rgb = AZUL

    derecha = tabla.cell(0, 1).paragraphs[0]
    derecha.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    titulo = derecha.add_run(ORGANIZACION)
    titulo.bold = True
    titulo.font.size = Pt(12)
    titulo.font.color.rgb = AZUL
    derecha.add_run("\n")
    sub = derecha.add_run(UNIDAD)
    sub.font.size = Pt(9)
    sub.font.color.rgb = AZUL

    # PIE. Una sola línea centrada y chica, en todas las páginas.
    pie = seccion.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    linea = pie.add_run(PIE)
    linea.font.size = Pt(8)
    linea.font.color.rgb = AZUL

    # El cuerpo queda vacío a propósito: `markdown_a_docx()` lo borra igual.
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    for nombre in ESTILOS_A_QUITAR:
        try:
            doc.styles[nombre].delete()
        except KeyError:
            pass

    return doc


def main():
    if DESTINO.exists() and "--sobrescribir" not in sys.argv:
        print(f"Ya existe {DESTINO}.")
        print("Pasale --sobrescribir si querés regenerarla.")
        return 1
    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    construir().save(str(DESTINO))
    print(f"Listo: {DESTINO}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
